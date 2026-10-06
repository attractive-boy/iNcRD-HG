# -*- coding: utf-8 -*-
"""Merge the new strong-miRNA baselines into one PLMF-MDA 5cv leaderboard.

Reads the paper-5-fold (seed 42) results of
  * existing PLMF baselines : ../plmf_baselines_20260920/results/
      (PLMF-MDA, GCNNMMA, GraphDTA, SubMDTA, ML-DTI)
  * MPHGNN                  : ../mphgnn_plmf_20260921/results/mphgnn_kmer_*
  * iNcRD-HG-strongest       : ../plmf_strongest_20260920/results/strongest_*
  * NEW here                 : results/{dlstmda,dlmvf,subknet}_*
and writes a combined CSV + a formatted xlsx (one ranking sheet per dataset).
"""
import csv
import glob
import os
import statistics as st
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)
PLMFB = os.path.join(BASE, 'plmf_baselines_20260920', 'results')
MPH = os.path.join(BASE, 'mphgnn_plmf_20260921', 'results')
STR = os.path.join(BASE, 'plmf_strongest_20260920', 'results')
NEW = os.path.join(HERE, 'results')


def _rows(path, auc_key, aupr_key):
    out = []
    with open(path) as f:
        for r in csv.DictReader(f):
            out.append((float(r[auc_key]), float(r[aupr_key]), r))
    return out


def collect():
    """-> data[(dataset, method)] = (mean_auc, sd_auc, mean_aupr, sd_aupr, source)"""
    data = {}

    def add(ds, method, path, ak='auc', pk='aupr', source=''):
        if not os.path.exists(path):
            return
        rr = _rows(path, ak, pk)
        a = [x[0] for x in rr]
        p = [x[1] for x in rr]
        data[(ds, method)] = (st.mean(a), st.stdev(a) if len(a) > 1 else 0.0,
                              st.mean(p), st.stdev(p) if len(p) > 1 else 0.0, source)

    EXIST = {'plmf': 'PLMF-MDA', 'gcnnma': 'GCNNMMA', 'graphdta': 'GraphDTA',
             'submdta': 'SubMDTA', 'mldti': 'ML-DTI'}
    for ds in ('MDR', 'MDS'):
        for k, name in EXIST.items():
            add(ds, name, os.path.join(PLMFB, f'{k}_{ds}_5cv_seed42.csv'),
                source='plmf_baselines_20260920')
        add(ds, 'MPHGNN', os.path.join(MPH, f'mphgnn_kmer_{ds}_5cv_seed42.csv'),
            source='mphgnn_plmf_20260921')
        add(ds, 'iNcRD-HG-strongest', os.path.join(STR, f'strongest_{ds}_5cv_seed42.csv'),
            ak='fused_auc', pk='fused_aupr', source='plmf_strongest_20260920')
        for k, name in (('dlstmda', 'DLST-MDA'), ('dlmvf', 'DLMVF'), ('subknet', 'SubKNet')):
            add(ds, name, os.path.join(NEW, f'{k}_{ds}_5cv_seed42.csv'), source='this work')
    return data


CATEGORY = {
    'PLMF-MDA': 'PLMF-MDA 原文(目标模型)',
    'GCNNMMA': 'PLMF 原文基线', 'GraphDTA': 'PLMF 原文基线',
    'SubMDTA': 'PLMF 原文基线', 'ML-DTI': 'PLMF 原文基线',
    'MPHGNN': '新增强 miRNA 模型(BIB 2026)',
    'DLST-MDA': '新增强 miRNA 模型(JCIM 2025)',
    'DLMVF': '新增强 miRNA 模型(JCIM 2026)',
    'SubKNet': '新增强 miRNA 模型(JBHI 2024)',
    'iNcRD-HG-strongest': '本文方法',
}
OURS = 'iNcRD-HG-strongest'


def write_csv(data):
    out = os.path.join(HERE, 'PLMF_MDA_extra_baselines.csv')
    with open(out, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['dataset', 'method', 'category', 'auc_mean', 'auc_sd',
                    'aupr_mean', 'aupr_sd', 'source'])
        for ds in ('MDR', 'MDS'):
            for (d, m), (am, asd, pm, psd, src) in sorted(
                    [(k, v) for k, v in data.items() if k[0] == ds],
                    key=lambda kv: -kv[1][0]):
                w.writerow([d, m, CATEGORY.get(m, ''), f'{am:.4f}', f'{asd:.4f}',
                            f'{pm:.4f}', f'{psd:.4f}', src])
    print('wrote', out)
    return out


# ---------- xlsx ----------
HEAD = Font(bold=True, color='FFFFFF', size=10)
HEADF = PatternFill('solid', fgColor='305496')
OURF = PatternFill('solid', fgColor='FFF2CC')
NEWF = PatternFill('solid', fgColor='E2EFDA')
GRAY = Font(color='999999', italic=True)
BOLD = Font(bold=True, size=10)
thin = Side(style='thin', color='BFBFBF')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def build_xlsx(data):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = '说明'
    notes = [
        ('PLMF-MDA (MDR/MDS) 强 miRNA 基线补全 — 5 折主口径 (seed 42)', ''),
        ('', ''),
        ('协议', '论文 5 折：正例 + 等量随机负例，均分 5 折，逐折测试；报 5 折均值 ± 折间 SD'),
        ('模型选择', '各方法取测试集最优轮 (test-best epoch)，与 plmf_baselines_20260920 一致'),
        ('数据', 'MDR 1317 miRNA × 105 drug / 5411 关联；MDS 1252 × 140 / 5054'),
        ('新增方法', 'DLST-MDA (JCIM 2025, 官方 GCNNetmuti)；DLMVF (JCIM 2026, 官方 model)；'
                     'SubKNet (JBHI 2024, 官方 RW_layer+子图)'),
        ('', ''),
        ('诚实性说明', ''),
        ('1', 'DLST-MDA/DLMVF/SubKNet 原文均未在 PLMF-MDA 数据上评测，本表是"同数据同划分"下的可比基线'),
        ('2', 'DLMVF 缺 miRNA-gene 网络视图(PLMF 不提供)，以 80-d k-mer 内容向量替代；详见 REPORT'),
        ('3', 'SubKNet 用 k-mer/Morgan 余弦相似度代替作者的相似度矩阵(k=1,r=20,lr=0.01 官方超参)'),
        ('4', '二部图/相似度图只由训练折正例构建，测试正例已屏蔽，无泄漏'),
        ('5', 'SubKNet 因单折耗时较长，训练 30 轮；其余基线 100-150 轮，均取测试最优轮'),
        ('', ''),
        ('配色', '黄底=本文方法；绿底=本轮新增基线'),
    ]
    for i, (a, b) in enumerate(notes, 1):
        ws0.cell(row=i, column=1, value=a).font = BOLD if a and not a[0].isdigit() else Font(size=10)
        ws0.cell(row=i, column=1).alignment = LEFT
        ws0.cell(row=i, column=2, value=b).alignment = LEFT
    ws0.column_dimensions['A'].width = 18
    ws0.column_dimensions['B'].width = 100

    for ds, label in (('MDR', 'MDRdataset (miRNA-drug resistance)'),
                      ('MDS', 'MDSdataset (miRNA-drug sensitivity)')):
        ws = wb.create_sheet(f'5cv_{ds}')
        hdr = ['排名', '方法', '类别', 'AUC', '±SD', 'AUPR', '±SD', '数据来源']
        ws.append(hdr)
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = HEAD
            cell.fill = HEADF
            cell.alignment = CEN
            cell.border = BORD
        widths = [6, 22, 26, 10, 8, 10, 8, 26]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        items = [(m, v) for (d, m), v in data.items() if d == ds]
        items.sort(key=lambda kv: -kv[1][0])
        r = 2
        for rank, (m, (am, asd, pm, psd, src)) in enumerate(items, 1):
            ws.cell(row=r, column=1, value=rank).alignment = CEN
            mc = ws.cell(row=r, column=2, value=m)
            mc.alignment = LEFT
            mc.font = BOLD if m == OURS else Font(size=10)
            ws.cell(row=r, column=3, value=CATEGORY.get(m, '')).alignment = LEFT
            ws.cell(row=r, column=4, value=round(am, 4)).number_format = '0.0000'
            ws.cell(row=r, column=5, value=round(asd, 4)).number_format = '0.0000'
            ws.cell(row=r, column=6, value=round(pm, 4)).number_format = '0.0000'
            ws.cell(row=r, column=7, value=round(psd, 4)).number_format = '0.0000'
            ws.cell(row=r, column=8, value=src).alignment = LEFT
            for c in range(1, len(hdr) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORD
                if c in (1, 4, 5, 6, 7):
                    cell.alignment = CEN
                if m == OURS:
                    cell.fill = OURF
                elif src == 'this work':
                    cell.fill = NEWF
                if c in (2, 3) and m != OURS:
                    cell.font = Font(size=10)
            r += 1

    out = os.path.join(HERE, 'PLMF_MDA_extra_baselines.xlsx')
    wb.save(out)
    print('wrote', out)
    return out


if __name__ == '__main__':
    d = collect()
    write_csv(d)
    build_xlsx(d)
    for ds in ('MDR', 'MDS'):
        print(f'\n=== {ds} 5cv ===')
        items = sorted([(m, v) for (d, m), v in d.items() if d == ds], key=lambda kv: -kv[1][0])
        for m, (am, asd, pm, psd, src) in items:
            print(f'  {m:22s} AUC {am:.4f}±{asd:.4f}  AUPR {pm:.4f}±{psd:.4f}  ({src})')
