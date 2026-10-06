# -*- coding: utf-8 -*-
"""Rank iNcRD-HG's strongest portable model vs the MSFFMDI baselines on ncDR / RNAInter."""
import os
import statistics as st

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)
BL = os.path.join(BASE, 'msffmdi_baselines_20260920', 'results')
OUT = os.path.join(HERE, 'MSFFMDI_strongest_ranking.xlsx')

BASELINE = {'lagcn': 'LAGCN', 'gcnnma': 'GCNNMMA', 'mfidma': 'MFIDMA',
            'ldaman': 'LDAMAN', 'dlstmda': 'DLST-MDA'}
PAPER_MSFFMDI = {'ncDR': (0.9483, 0.9461, 0.8834), 'RNAInter': (0.9216, 0.9163, 0.8524)}

HEAD = Font(bold=True, color='FFFFFF', size=10)
HEADF = PatternFill('solid', fgColor='305496')
OURF = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='BFBFBF')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center')


def strongest(dataset):
    d = pd.read_csv(os.path.join(HERE, 'results', f'strongest_{dataset}.csv'))
    return d


def baseline(dataset, key):
    fp = os.path.join(BL, f'{key}_{dataset}_seed42.csv')
    return pd.read_csv(fp) if os.path.exists(fp) else None


def board(dataset):
    e = []
    d = strongest(dataset)
    e.append(dict(method='iNcRD-HG-strongest (单模型 TabICLv2)', auc=d.ref_auc.mean(),
                  aupr=d.ref_aupr.mean(), acc=d.acc.mean(), sen=d.sen.mean(),
                  spec=d.spec.mean(), n=len(d), note='frozen TabICLv2 图上下文参考（无融合）',
                  ours=True))
    e.append(dict(method='iNcRD-HG-strongest (+anchor 融合)', auc=d.fused_auc.mean(),
                  aupr=d.fused_aupr.mean(), acc=d.acc.mean(), sen=d.sen.mean(),
                  spec=d.spec.mean(), n=len(d), note='单模型 + anchor-shrinkage (w=0.3)',
                  ours=True))
    for k, name in BASELINE.items():
        b = baseline(dataset, k)
        if b is None:
            continue
        e.append(dict(method=name, auc=b.auc.mean(), aupr=b.aupr.mean(), acc=b.acc.mean(),
                      sen=b.sen.mean(), spec=b.spec.mean(), n=len(b), note='复现基线', ours=False))
    pa, pp, pacc = PAPER_MSFFMDI[dataset]
    e.append(dict(method='MSFFMDI (论文)', auc=pa, aupr=pp, acc=pacc, sen=None, spec=None,
                  n=None, note='论文报告值', ours=False))
    e.sort(key=lambda r: (-r['auc'], -r['aupr']))
    for i, r in enumerate(e, 1):
        r['rank'] = i
    return e


def main():
    wb = Workbook(); ws0 = wb.active; ws0.title = '说明'
    notes = [
        ('MSFFMDI 数据：本文最强模型排名', ''),
        ('本文方法', 'iNcRD-HG 最强框架的可移植子集：frozen TabICLv2 图上下文参考 + anchor-shrinkage'),
        ('移植范围', '仅图上下文 33 列（关联图 + k-mer/Morgan 相似度）；证据档/DepMap 档因 ncDR/RNAInter 无数据未启用'),
        ('协议', '与 msffmdi_baselines_20260920 完全相同的 clean 5 折 CV（去重正例 + 等量负例，seed42）'),
        ('对照', 'LAGCN/GCNNMMA/MFIDMA/LDAMAN/DLST-MDA 复现值 + 论文 MSFFMDI 报告值'),
    ]
    for i, (a, b) in enumerate(notes, 1):
        ws0.cell(i, 1, a).font = Font(bold=True, size=10); ws0.cell(i, 1).alignment = LEFT
        ws0.cell(i, 2, b).alignment = LEFT
    ws0.column_dimensions['A'].width = 14; ws0.column_dimensions['B'].width = 95

    for ds in ('ncDR', 'RNAInter'):
        rows = board(ds)
        ws = wb.create_sheet(ds)
        hdr = ['排名', '方法', 'AUC', 'AUPR', 'ACC', 'Sen', 'Spec', '运行数', '备注']
        ws.append([f'{ds}dataset — 本文最强模型 vs MSFFMDI 基线（5 折均值）'])
        ws.cell(1, 1).font = Font(bold=True, size=11)
        ws.append(hdr)
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(2, c); cell.font = HEAD; cell.fill = HEADF; cell.alignment = CEN; cell.border = BORD
        ws.freeze_panes = 'A3'
        r = 3
        for e in rows:
            vals = [e['rank'], e['method'], e['auc'], e['aupr'], e['acc'], e['sen'], e['spec'],
                    e['n'], e['note']]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c)
                if isinstance(v, float):
                    cell.value = round(v, 4); cell.number_format = '0.0000'
                else:
                    cell.value = v
                cell.alignment = LEFT if c in (2, 9) else CEN
                cell.border = BORD
                if e['ours']:
                    cell.fill = OURF; cell.font = Font(bold=True, size=10)
            r += 1
        for i, w in enumerate([6, 30, 9, 9, 9, 9, 9, 7, 34], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        print(f'\n=== {ds} ===')
        for e in rows:
            print(f"{e['rank']:>2} {e['method']:28s} AUC {e['auc']:.4f}  AUPR {e['aupr']:.4f}  ACC {e['acc']:.4f}")
    wb.save(OUT)
    print('\nwrote', OUT)


if __name__ == '__main__':
    main()
