# -*- coding: utf-8 -*-
"""Build the Excel for the PLMF-MDA GNN-architecture ablation (plmf_gnnarch_20260923)
from the per-fold CSVs in results/.  Metric = valsel (hold-out-selected epoch) AUC/AUPR,
mean over the 5 paper folds, seed 42; deltas are paired per-fold against the reference tag.
"""
import os, glob
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(HERE, 'results')
OUT = os.path.join(HERE, '消融实验_GNN架构.xlsx')
DS = ('MDR', 'MDS')


def load(tag, ds):
    f = os.path.join(RES, f'{tag}_{ds}_5cv_seed42.csv')
    return pd.read_csv(f).set_index('fold') if os.path.exists(f) else None


def stats(tag, ref):
    """mean valsel AUC/AUPR over folds for MDR/MDS + paired delta vs ref."""
    out = {}
    for ds in DS:
        d = load(tag, ds)
        if d is None:
            out[ds] = None; continue
        a, p = d.valsel_auc.mean(), d.valsel_aupr.mean()
        da = dp = wins = None
        r = load(ref, ds) if ref else None
        if r is not None:
            j = d.join(r, rsuffix='_r', how='inner')
            da = (j.valsel_auc - j.valsel_auc_r).mean()
            dp = (j.valsel_aupr - j.valsel_aupr_r).mean()
            wins = int(((j.valsel_auc - j.valsel_auc_r) > 0).sum())
        out[ds] = dict(auc=a, aupr=p, da=da, dp=dp, wins=wins, n=len(d))
    return out


# (阶段, 名称, 命令行, tag, 参考tag)
LADDER = [
    ('0', 'plain iNcRD-HG（GeneralConv×2 + MLP）', '', 'plain', None),
    ('1', '+ LQO 留出查询边消息传递', '--lqo 5', 'lqo', 'plain'),
    ('2', '+ RT-NCN 关系型神经共同邻居配对解码器', '--lqo 5 --cnt --ncn', 'pair', 'lqo'),
    ('3', '+ 溯源超边配对通道（study / 语境 token）', '--lqo 5 --cnt --ncn --prov', 'pair_prov', 'pair'),
    ('4', '+ 旁信息关系配对通道（HMDD 疾病 / 家族 / 基因组簇）', '--lqo 5 --cnt --ncn --prov --side', 'full', 'pair_prov'),
    ('对照', '溯源记录在配对间置换（full_provS）', '--lqo 5 --cnt --ncn --prov --side --prov_shuffle', 'full_provS', 'full'),
    ('对照', '旁信息在 miRNA 间置换（full_sideS）', '--lqo 5 --cnt --ncn --prov --side --side_shuffle', 'full_sideS', 'full'),
]
M1M2 = [
    ('lqo（仅 LQO）', '--lqo 5', 'lqo', None),
    ('lqo + cnt（仅计数）', '--lqo 5 --cnt', 'lqo_cnt', 'lqo'),
    ('lqo + ncn（仅隐状态）', '--lqo 5 --ncn', 'lqo_ncn', 'lqo'),
    ('lqo + cnt + ncn = pair', '--lqo 5 --cnt --ncn', 'lqo_pair', 'lqo'),
]
M3 = [
    ('pair + prov（study + token，图 + 配对通道）', '--lqo 5 --cnt --ncn --prov', 'pair_prov', 'pair'),
    ('仅 study 超边', '--lqo 5 --cnt --ncn --prov --prov_groups study', 'pair_prov_study', 'pair'),
    ('仅 token 超边', '--lqo 5 --cnt --ncn --prov --prov_groups ctxtok', 'pair_prov_tok', 'pair'),
    ('仅作为图节点（--prov_pair 0，无配对通道）', '--lqo 5 --cnt --ncn --prov --prov_pair 0', 'pair_prov_g', 'pair'),
    ('记录置换对照（--prov_shuffle）', '--lqo 5 --cnt --ncn --prov --prov_shuffle', 'pair_provS', 'pair'),
]
M4 = [
    ('pair + side（三组，图 + 配对通道）', '--lqo 5 --cnt --ncn --side', 'pair_side', 'pair'),
    ('仅 HMDD 疾病', '--lqo 5 --cnt --ncn --side --side_groups disease', 'pair_side_dz', 'pair'),
    ('仅家族', '--lqo 5 --cnt --ncn --side --side_groups family', 'pair_side_fam', 'pair'),
    ('仅基因组簇', '--lqo 5 --cnt --ncn --side --side_groups cluster', 'pair_side_clu', 'pair'),
    ('旁信息置换对照（--side_shuffle）', '--lqo 5 --cnt --ncn --side --side_shuffle', 'pair_sideS', 'pair'),
]
NOLQO = [
    ('plain（无模块、无 LQO）', '', 'plain', None),
    ('+ cnt', '--cnt', 'cnt_nolqo', 'plain'),
    ('+ ncn', '--ncn', 'ncn_nolqo', 'plain'),
    ('+ RT-NCN (cnt+ncn)', '--cnt --ncn', 'pair_nolqo', 'plain'),
    ('+ RT-NCN + prov', '--cnt --ncn --prov', 'pair_prov_nolqo', 'plain'),
    ('+ RT-NCN + prov + side', '--cnt --ncn --prov --side', 'full_nolqo', 'plain'),
]

HFILL = PatternFill('solid', fgColor='1F4E78')
HFONT = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
wb.remove(wb.active)


def add_sheet(name, rows, title):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ['阶段', '模块 / 变体', '命令行参数',
           'MDR AUC', 'MDR AUPR', 'ΔAUC vs 参考', 'ΔAUPR vs 参考', '折正/5',
           'MDS AUC', 'MDS AUPR', 'ΔAUC vs 参考', 'ΔAUPR vs 参考', '折正/5']
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=3, column=c); cell.fill = HFILL; cell.font = HFONT
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for r in rows:
        if len(r) == 4:
            r = ('',) + r
        stage, name, cmd, tag, ref = r
        s = stats(tag, ref)
        row = [stage, name, cmd]
        for ds in DS:
            v = s[ds]
            if v is None:
                row += ['', '', '', '', '']
            else:
                row += [round(v['auc'], 4), round(v['aupr'], 4),
                        (f"{v['da']:+.4f}" if v['da'] is not None else '—'),
                        (f"{v['dp']:+.4f}" if v['dp'] is not None else '—'),
                        (f"{v['wins']}/{v['n']}" if v['wins'] is not None else '—')]
        ws.append(row)
    widths = [7, 40, 42, 10, 10, 13, 13, 8, 10, 10, 13, 13, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=4):
        row[1].alignment = WRAP; row[2].alignment = WRAP
    return ws


add_sheet('0_阶梯消融', LADDER,
          'PLMF-MDA GNN 架构消融 · 阶梯（论文 5 折 seed 42，valsel AUC/AUPR，5 折均值；Δ 为对上一级同折配对差）')
add_sheet('1_M1-M2内部', M1M2,
          'M1 LQO / M2 RT-NCN 内部消融（Δ 对 lqo）')
add_sheet('2_M3溯源内部', M3,
          'M3 溯源超边配对通道内部消融（Δ 对 pair）')
add_sheet('3_M4旁信息内部', M4,
          'M4 旁信息关系配对通道内部消融（Δ 对 pair）')
add_sheet('4_无LQO对照', NOLQO,
          '不开 LQO 时各模块的效果（Δ 对 plain）——证明 LQO 是前置条件')

# 说明
ws = wb.create_sheet('说明')
notes = [
    ['消融实验_GNN架构.xlsx 说明', ''],
    ['', ''],
    ['来源', 'improve/research/plmf_gnnarch_20260923/（报告 REPORT_GNNARCH.md；逐折数据 results/*.csv）'],
    ['模型', 'iNcRD-HG GNN 本体：HeteroConv 逐关系实例化 GeneralConv(-1,-1,256)×2 + MLP 边解码器'],
    ['数据 / 划分', 'PLMF-MDA MDR / MDS；论文 5 折，seed 42（get_split，与基线复现逐位相同）'],
    ['指标口径', 'valsel：从训练对分层留出 10% 选 epoch，取该 epoch 的测试 AUC/AUPR；5 折均值'],
    ['Δ 与折正', 'Δ = 同折配对差均值；折正 = 5 折中 ΔAUC>0 的折数'],
    ['K', 'LQO 组数 K=5；不开 LQO 时 K=1（消息图=全部训练正例）'],
    ['对照', 'full_provS / full_sideS 分别把溯源记录在正例间、旁信息行在 miRNA 间置换'],
    ['无 LQO 对照', 'cnt_nolqo / ncn_nolqo / pair_nolqo / pair_prov_nolqo / full_nolqo 为本次新增，命令行同对应变体但不加 --lqo'],
    ['复现', 'run_grid.sh grid_ladder.txt / grid_v2.txt；无 LQO 对照 run_nolqo_experiment.sh'],
]
for r in notes:
    ws.append(r)
ws['A1'].font = Font(bold=True, size=12)
ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 110
for row in ws.iter_rows():
    row[1].alignment = WRAP

wb.save(OUT)
print('written', OUT)
