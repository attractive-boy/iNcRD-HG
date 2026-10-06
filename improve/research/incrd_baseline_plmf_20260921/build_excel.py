# -*- coding: utf-8 -*-
"""Write iNcRD-HG_baseline_on_PLMF-MDA.xlsx: README / 4 boards / variant deltas / per-run rows."""
import glob, os, statistics as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
import sys
sys.path.insert(0, HERE)
import summarize as S  # noqa: E402

HDR = PatternFill('solid', fgColor='DDEBF7'); BOLD = Font(bold=True); OURS = PatternFill('solid', fgColor='FFF2CC')


def sheet(wb, name, header, rows, widths=None, hl=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font = BOLD; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
        if hl and hl(r):
            for c in ws[ws.max_row]:
                c.fill = OURS
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(60, len(str(h)) + 4)))
    ws.freeze_panes = 'A2'
    return ws


def main():
    wb = Workbook(); wb.remove(wb.active)
    readme = [
        ['iNcRD-HG plain baseline (GeneralConv x2 + MLP edge decoder, predict_association.py) on PLMF-MDA MDR/MDS'],
        [''],
        ['Boards', 'MDR/MDS x paper 5-fold (seed 42, fold mean) and released-code 80/20 (seeds 0-4, mean); splits bit-identical to plmf_baselines_20260920'],
        ['Main metric column', 'test-best epoch (same optimistic convention as the PLMF baseline board); valsel = epoch chosen on a 10% training hold-out; refit = retrain on all training pairs for the val-chosen number of epochs, final-epoch test metric (leakage-free headline)'],
        ['Variants', 'see sheet "variants": tag -> extra CLI flags of run_v3.py; plain = original recipe with k-mer/ChemBERTa node features and top-10 cosine kNN similarity edges'],
        ['ctx33', '33-column label-free graph-context pair features (dlstmda_strongest_20260920/features.py::GraphPairEncoder, the portable tier of iNcRD-HG-strongest) computed from TRAINING positives only, 5-fold cross-fitted for training pairs, MLP-projected and concatenated to the decoder input'],
        ['ctx33_shuf', 'control: same 33 columns with rows permuted across pairs (destroys the pair information, keeps the marginal distribution)'],
        ['Reference rows', 'PLMF-MDA + 4 baselines from plmf_baselines_20260920; MPHGNN from mphgnn_plmf_20260921 (batch-scored test-best and val-selected isolated-pair); iNcRD-HG-strongest (frozen TabICLv2 graph-context + anchor) from plmf_strongest_20260920'],
    ]
    ws = wb.create_sheet('README')
    for r in readme:
        ws.append(r)
    ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 150
    T = S.tags()
    for ds, split in S.BOARDS:
        rows = []
        for t in T:
            r = S.ours(t, ds, split)
            if r:
                fs = sorted(glob.glob(os.path.join(HERE, 'results', f'incrd_{t}_{ds}_{split}_seed*.csv')))
                d = pd.concat([pd.read_csv(f) for f in fs])
                rows.append([f'iNcRD-HG[{t}]', r['auc'], r['auc_sd'], r['aupr'], r['aupr_sd'], r.get('last_auc'), r.get('last_aupr'),
                             d.valsel_auc.mean() if 'valsel_auc' in d else None, d.valsel_aupr.mean() if 'valsel_aupr' in d else None,
                             d.refit_auc.mean() if 'refit_auc' in d else None, d.refit_aupr.mean() if 'refit_aupr' in d else None,
                             d.best_epoch.mean(), r['n']])
        for m, r in S.refs(ds, split).items():
            rows.append([m, r['auc'], r['auc_sd'], r['aupr'], r['aupr_sd'], None, None, None, None, None, None, None, r['n']])
        rows.sort(key=lambda x: -x[3])
        rows = [[i + 1] + r for i, r in enumerate(rows)]
        sheet(wb, f'{ds}_{split}', ['rank', 'method', 'AUC (test-best)', 'AUC sd', 'AUPR (test-best)', 'AUPR sd', 'last AUC', 'last AUPR',
                                    'valsel AUC', 'valsel AUPR', 'refit AUC', 'refit AUPR', 'mean best epoch', 'n'],
              rows, {'method': 70}, hl=lambda r: str(r[1]).startswith('iNcRD-HG['))
    # variants
    base = {b: S.ours('plain', *b) for b in S.BOARDS}
    rows = []
    for t in T:
        row = [t]; da, dp = [], []
        for b in S.BOARDS:
            r = S.ours(t, *b); p = base[b]
            if r and p:
                row += [r['auc'], r['aupr'], r['auc'] - p['auc'], r['aupr'] - p['aupr']]; da.append(r['auc'] - p['auc']); dp.append(r['aupr'] - p['aupr'])
            else:
                row += [None] * 4
        row += [st.mean(da) if da else None, st.mean(dp) if dp else None]
        rows.append(row)
    rows.sort(key=lambda x: -(x[-1] or -9))
    hdr = ['tag'] + [f'{d} {s} {m}' for d, s in S.BOARDS for m in ('AUC', 'AUPR', 'dAUC', 'dAUPR')] + ['mean dAUC', 'mean dAUPR']
    sheet(wb, 'variants', hdr, rows)
    # per run
    fs = sorted(glob.glob(os.path.join(HERE, 'results', 'incrd_*_seed*.csv')))
    d = pd.concat([pd.read_csv(f).assign(tag=os.path.basename(f)[6:].rsplit('_', 3)[0]) for f in fs], ignore_index=True)
    cols = ['tag'] + [c for c in d.columns if c != 'tag']
    sheet(wb, 'runs', cols, d[cols].where(pd.notnull(d[cols]), None).values.tolist())
    out = os.path.join(HERE, 'iNcRD-HG_baseline_on_PLMF-MDA.xlsx')
    wb.save(out); print('wrote', out)


if __name__ == '__main__':
    main()
