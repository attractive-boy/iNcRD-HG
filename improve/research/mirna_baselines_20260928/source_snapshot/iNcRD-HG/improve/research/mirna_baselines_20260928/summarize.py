"""Build the three-dataset table from actual verified run files, never paper numbers."""
import argparse,csv,json,math,statistics,unicodedata
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from bench_data import HERE,ROOT

METRICS=['auc','aupr','acc','precision','recall','specificity','f1','mcc']

def display_width(value):
    return sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in str(value or ''))

def main(require_complete=False):
    reg=json.loads((HERE/'registry.json').read_text()); groups=defaultdict(list);runs=[]
    expected_runs=len(reg)*13; expected_cells=len(reg)*3
    for path in (HERE/'results').glob('*.json'):
        if path.name.endswith('_curves.json'):continue
        r=json.loads(path.read_text());runs.append(r);groups[r['model'],r['dataset']].append(r)
    summary=[];matrix=[];pending=[]
    for model in reg:
        row={'模型':model['name'],'类别':model['category']}
        for ds in ('MiDrug','MDR','MDS'):
            rs=groups[model['id'],ds];nexp=3 if ds=='MiDrug' else 5
            complete=len(rs)==nexp
            if not complete:pending.append(f'{model["name"]} / {ds}: {len(rs)}/{nexp}')
            rec={'model':model['id'],'dataset':ds,'n':len(rs),'expected':nexp,'status':'complete' if complete else 'in_progress'}
            for k in METRICS:
                values=[r[k] for r in rs]
                rec[k+'_mean']=statistics.mean(values) if complete else None
                rec[k+'_sd']=statistics.stdev(values) if complete else None
            summary.append(rec)
            for k in ('auc','aupr'):
                row[ds+' '+k.upper()]=(f'{rec[k+"_mean"]:.4f} ± {rec[k+"_sd"]:.4f}' if complete else f'进行中 {len(rs)}/{nexp}')
        matrix.append(row)
    if require_complete and pending:raise RuntimeError('Incomplete:\n'+'\n'.join(pending))
    with (HERE/'summary.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(summary[0]));w.writeheader();w.writerows(summary)
    wb=Workbook();ws=wb.active;ws.title='三数据集对照'
    ws.append(list(matrix[0]));
    for row in matrix:ws.append(list(row.values()))
    ws=wb.create_sheet('全部指标汇总')
    keys=list(summary[0]);ws.append(keys)
    for r in summary:ws.append([r[k] for k in keys])
    ws=wb.create_sheet('逐次运行')
    keys=['model','dataset','fold','seed','auc','aupr','acc','precision','recall','specificity','f1','mcc',
          'best_epoch','epochs_run','val_auc','elapsed_seconds','selection','graph','threshold','split_sha256','adapter_note']
    ws.append(keys)
    for r in sorted(runs,key=lambda r:(r['model'],r['dataset'],r['fold'],r['seed'])):ws.append([r.get(k) for k in keys])
    ws=wb.create_sheet('模型与论文来源');keys=list(reg[0]);ws.append(keys)
    for r in reg:ws.append([r[k] for k in keys])
    ws=wb.create_sheet('运行超参数');ws.append(['模型','数据集','折','种子','配置 JSON'])
    for r in runs:ws.append([r['model'],r['dataset'],r['fold'],r['seed'],json.dumps(r['config'],ensure_ascii=False)])
    manifest=json.loads((HERE/'data/manifest.json').read_text())
    ws=wb.create_sheet('数据与划分');ws.append(['数据集','miRNA 数','药物数','折','训练总数','训练正例','验证总数','测试总数','测试正例','ID 对重叠','内容相同 train-test 对数','划分 SHA256'])
    for ds,info in manifest['datasets'].items():
        for f,r in info['folds'].items():ws.append([ds,info['n_mirna'],info['n_drug'],f,r['train']['n'],r['train']['positive'],r['val']['n'],r['test']['n'],r['test']['positive'],0,r['content_pair_train_test_overlap'],r['sha256']])
    notes=[
     ('完成状态',f'{len(runs)}/{expected_runs} 次正式运行；{expected_cells-len(pending)}/{expected_cells} 个完整模型×数据集组合。短跑 smoke 不计入。'),
     ('MiDrug','保留原有 train/val/test 成员；只删除训练中 2 条同标签重复行。605 miRNA × 216 药物；正负 1:2。3 个模型种子 42/43/44，标准差反映初始化波动。'),
     ('MDR/MDS','作者公开代码为单次 80/20；本项目按论文描述构建 seed=42 的外层 5 折，与前续基线逐行相同，并非作者发布的折索引。每个训练折中分层取 10% 验证。正负 1:1。只做一套外层划分，标准差为折间波动，不是 5 次独立重复。'),
     ('选模型','仅以验证 AUROC 选 epoch；固定 0.5 分类阈值；测试标签只在选定模型后用于一次评估。未做嵌套超参搜索。'),
     ('图与特征','目标 miRNA–drug 关联图/GIP 只使用训练正例；验证正例与测试正例均不进入消息图。冻结内容特征可使用已知实体全集。原始 iNcRD-HG 在 MiDrug 上额外保留发布的 lncRNA 辅助关系与表达谱，方法间输入信息并非完全相同。'),
     ('成熟序列','MiDrug 605 条成熟序列已按名称补全；604 条不同序列。以 PLMF 元数据表为历史名称来源，miRBase 为补充，miR-302f 的 arm 更名单独记录。未把前体序列头部误作成熟序列。'),
     ('内容重合','作者实体 ID 未合并；按相同成熟序列+原始 SMILES 字符串检查，MiDrug 有 4 个、MDR 每折 2–4 个、MDS 每折 6–9 个 train-test 内容相同对。ID 对无重叠。本轮是常规转导式关联预测，不是内容去重或冷启动。'),
     ('AP 定义','AUPR 使用 sklearn average_precision_score（AP），不是梯形积分 PR-AUC；不同正负比例下不直接比较三个数据集的 AP 高低。'),
     ('负例语义','使用公开 benchmark 未关联/负例池；这些不是全部经过生物实验验证的真阴性，属于闭世界采样基准。'),
     ('与旧结果比较','旧脚本多按 test-best 选轮次，部分还在测试图加入验证正例。本次严格 train-only/validation-select，不能把旧分数直接拼入主表。'),
     ('模型来源差异','PLMF/DLST/DLMVF/SubKNet 等保留作者主要网络并适配数据；GCMDR/AMMGC/LAGCN 是本地结构重实现。AMMGC 本地版联合训练四视图，并非原代码分阶段训练。'),
     ('不完整方法','SubMDTA 只有监督预测结构，缺原文子结构预训练；DLMVF 缺原始 miRNA–基因图，使用 k-mer 替代视图；JMSSMMA 使用本任务训练负例替代原作者任意节点类型负例。必须随分数保留这些标记。'),
     ('PLMF 适配修复','原函数截节点到 100 却保留全部键；MiDrug 的一个 551 原子药物因此越界/跨图。现保留大分子的全部原子与键，小分子保持原 100-node 填充；MDR/MDS 最大 96 原子，处理不变。已检查全部药物 batch 的边不跨分子。'),
     ('原始 iNcRD-HG','锁定最初代码提交 eb8de624a731a4c076143bf2f9c91e539affee43 的 Model/GNNEncoder/EdgeDecoder，GeneralConv×2、hidden=256、关系 mean、三层 MLP、全批量 RMSprop(lr=1e-4, alpha=.99)，训练 300 epochs。MiDrug 恢复原始 lncRNA/miRNA/drug 三类节点及辅助关系，去除后加的 circRNA；使用发布表达谱/ChemBERTa 特征。'),
     ('原始版跨数据集适配','MDR/MDS 缺原始表达谱与 lncRNA 辅助图：使用 80 维 2+3-mer 和原版 ChemBERTa-zinc-base-v1 768 维特征，行 min-max、双向余弦 kNN(10)，保留原网络及训练配方。原始版一行的 PLMF adapted 标记必须保留。目标边用共享训练正例，验证集选轮；不是原脚本逐轮看测试/测试图加入验证边的旧分数。'),
     ('统计范围','表中均值±样本标准差；没有把 3 个种子和 5 折混作相同重复数，也没有跨数据集平均排名；没有据这些重实现分数推翻原文结论。'),
     ('复核','verify.py 核对每条结果的预测、指标、checkpoint hash、数据 hash 和验证集最优 epoch。literature/verified_references.json 与 registry.json 保留论文出处。'),
    ]
    ws=wb.create_sheet('说明与限制');ws.append(['项目','说明'])
    for row in notes:ws.append(row)
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='24476B');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
        for col in ws.columns:
            letter=get_column_letter(col[0].column)
            width=min(65,max(14,max(display_width(c.value) for c in col[:30])+2))
            ws.column_dimensions[letter].width=width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,float):cell.number_format='0.0000'
            lines=max(math.ceil(display_width(c.value)/max(1,ws.column_dimensions[c.column_letter].width-2)) for c in row)
            ws.row_dimensions[row[0].row].height=min(180,max(20,15*lines))
        ws.row_dimensions[1].height=32
    target=ROOT/'三数据集_miRNA基线实验_20260928.xlsx';wb.save(target)
    md=['# 三数据集 miRNA 基线实验（2026-09-28）','',f'正式运行完成 **{len(runs)}/{expected_runs}**；完整组合 **{expected_cells-len(pending)}/{expected_cells}**。','',
        '每个数据集使用同一套划分、负例和验证集选轮规则。表中的 adapted/reimpl./supervised-only 标记是方法定义的一部分。','',
        '| '+' | '.join(matrix[0])+' |','| '+' | '.join(['---']*len(matrix[0]))+' |']
    md+=['| '+' | '.join(str(v) for v in row.values())+' |' for row in matrix]
    md+=['','## 协议与解释','']+['- **'+k+'**：'+v for k,v in notes]
    md+=['','## 候选模型与论文来源','','| 模型 | 类别 | 年份 | 论文 DOI | 实现 |','| --- | --- | --- | --- | --- |']
    md+=['| '+r['name']+' | '+r['category']+' | '+r['year']+' | '+('['+r['doi']+'](https://doi.org/'+r['doi']+')' if r['doi'] else '[原始源码]('+r['code']+')')+' | '+r['implementation']+' |' for r in reg]
    md+=['','## 产物','',f'- Excel：`{target}`。','- `summary.csv`：数值均值和标准差。',
         '- `results/*_predictions.npz`：逐样本 ID、标签、概率；`*.pt`：验证集选定的模型。',
         '- `results/*_curves.json`：训练损失和验证指标；没有逐轮测试分数。',
         '- `data/manifest.json`、`registry.json`、`literature/verified_references.json`：数据、方法和引用溯源。',
         '- `verification.json`：完整性和结果重算核查。','- 运行方式见 `README.md`。','']
    (HERE/'REPORT.md').write_text('\n'.join(md))
    print(target);print('Completed',len(runs),'/',expected_runs,'; complete cells',expected_cells-len(pending),'/',expected_cells)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--require-complete',action='store_true');a=p.parse_args();main(a.require_complete)
