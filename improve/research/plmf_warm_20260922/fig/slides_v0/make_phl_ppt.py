# -*- coding: utf-8 -*-
"""Build PHL architecture figure (CN + EN) and degree-stratified diagnostics as a .pptx.

Numbers come from ../REPORT_WARM_PHL.md (§0 table, §1.1, §4.3).
Palette: dataviz reference palette (slots 1-3 validated all-pairs, light mode).
"""
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Microsoft YaHei"
INK, INK2, MUTED = "0B0B0B", "52514E", "898781"
GRID, BASE, SURF, TILE = "E1E0D9", "C3C2B7", "FCFCFB", "F4F3EF"
GOOD = "006300"
BLUE, ORANGE, AQUA = "2A78D6", "EB6834", "1BAF7A"
BLUE_T, ORANGE_T, AQUA_T = "DFEBF9", "FCE8E1", "DDF3EB"


def rgb(h):
    return RGBColor.from_string(h)


# ----------------------------------------------------------------- text helpers
def _set_ea(rPr, name):
    latin = rPr.find(qn("a:latin"))
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.Element(qn("a:ea"))
        if latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    ea.set("typeface", name)


def set_font(run, size, bold=False, color=INK, italic=False, name=FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = name
    _set_ea(run._r.get_or_add_rPr(), name)


def fill_text(tf, lines, size=9, color=INK, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, margins=(0.06, 0.04, 0.06, 0.04), space_after=1.5,
              line_spacing=1.05, wrap=True):
    """lines: list of str | list-of-runs; a run is (text, {size,bold,color,italic})."""
    tf.word_wrap = wrap
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = anchor
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if isinstance(line, str):
            runs = [(line, {})]
        elif isinstance(line, tuple) and len(line) == 2 and isinstance(line[1], dict):
            runs = [line]
        else:
            runs = line
        for text, kw in runs:
            r = p.add_run()
            r.text = text
            set_font(r, kw.get("size", size), kw.get("bold", bold), kw.get("color", color),
                     kw.get("italic", False))
    return tf


def add_text(slide, x, y, w, h, lines, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(tb.text_frame, lines, **kw)
    return tb


# ---------------------------------------------------------------- shape helpers
def add_box(slide, x, y, w, h, fill=None, line=None, line_w=0.75, radius=0.10,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, alpha=None, dash=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 1000)))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            ln = shp.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", dash)
    else:
        shp.line.fill.background()
    return shp


def add_line(slide, x1, y1, x2, y2, color=INK2, width=1.0, head=False, dash=None,
             kind=MSO_CONNECTOR.STRAIGHT):
    c = slide.shapes.add_connector(kind, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", dash)
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def node_circle(slide, cx, cy, r, fill, line=None):
    return add_box(slide, cx - r, cy - r, 2 * r, 2 * r, fill=fill, line=line, line_w=0.75,
                   shape=MSO_SHAPE.OVAL)


def node_square(slide, cx, cy, r, fill, line=None):
    return add_box(slide, cx - r, cy - r, 2 * r, 2 * r, fill=fill, line=line, line_w=0.75,
                   radius=0.2)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(SURF)
    return s


# ------------------------------------------------------------------- strings
S = {
    "cn": dict(
        title="PHL：溯源超图启发式学习器（Provenance-Hypergraph Heuristic Learner）",
        subtitle="在 PLMF-MDA 支持图上叠加策展溯源超图与 miRNA 旁信息关系，确定性计算 128 列启发式，"
                 "接一个逻辑回归读头（约 130 个参数）。单模型、线性、无泄漏。",
        cols=["① 数据源", "② 多关系超图（仅用 support 正例构建）", "③ 启发式特征（128 列）", "④ 线性读头"],
        rows=[
            dict(src=("PLMF-MDA 训练正例", "miRNA–药物关联矩阵（MDR / MDS），仅训练折的正例"),
                 g_title="支持图 + kNN 内容图",
                 g_cap=["二部支持图：仅训练正例", "kNN 内容图：miRNA / 药物相似度", "实线 = 支持边；虚线 = kNN 边"],
                 f_title="ctx (33) + diff (52)",
                 f_body=["度、共邻、path3/5、SVD16、属性支撑", "PPR / Katz / path7 及行、列归一化", "（原 iNcRD-HG 特征库）"]),
            dict(src=("iNcRD-HG 溯源记录", "每条关联的 PMID，以及报道它的细胞系 / 表型 / 疾病语境"),
                 g_title="溯源超图",
                 g_cap=["研究超边：同一 PMID 报道的全部 (m, d)", "语境超边：细胞系 / 表型 / 疾病 token",
                        "共研究图 S = M·Mᵀ"],
                 f_title="prov (24)",
                 f_body=["PMID 桥接 |P_m ∩ P_d|、规模加权、Jaccard (6)", "共研究传递 (S_r A)、(A S_d) (6)",
                         "语境 token 交集 (4) · 记录统计 (8)"]),
            dict(src=("miRNA 外部知识（标签无关）", "HMDD v4 疾病谱 · 名称 / seed 家族 · 另一表型矩阵中的度"),
                 g_title="miRNA 旁信息关系",
                 g_cap=["miRNA → 家族 / 疾病 / 跨表型度", "标签无关，专门给零度、低度 miRNA 补信息"],
                 f_title="mprior (19)",
                 f_body=["名称家族 (5) · seed 7-mer 家族 (5)", "HMDD 疾病桥 Σ H[m,dz]·(HᵀA)[dz,d] (7)",
                         "另一表型的度 xdeg (2)"]),
        ],
        g1_labels=("miRNA", "drug"),
        g2_labels=("PMID 研究超边", "桥接 m₁–d₂"),
        g3_labels=("miRNA", ["名称 / seed 家族", "HMDD 疾病谱", "另一表型的度"]),
        note=[("无泄漏交叉拟合", dict(bold=True, size=9.5)),
              "训练行：K=10 内折，query 对绝不进入自己的 support",
              "测试行：用全部训练正例作 support",
              "所有启发式只读 support 正例"],
        strip_label="拼接 128 列",
        strip=["ctx+diff 85", "prov 24", "mprior 19"],
        lr=[("z = signed log1p + 标准化", dict(size=9)),
            ("score(m, d) = w·z + b", dict(size=10.5, bold=True)),
            ("逻辑回归，C = 0.3，约 130 个参数", dict(size=8.5, color=INK2))],
        out="score(m, d)",
        tiles=[("MDR · AUC / AUPR（论文 5 折，4 seeds）", "0.9591 / 0.9565", "+0.0062 / +0.0071 vs 基线 0.9529 / 0.9494（4/4 seeds）"),
               ("MDS · AUC / AUPR（论文 5 折，4 seeds）", "0.9607 / 0.9587", "+0.0047 / +0.0051 vs 基线 0.9560 / 0.9536（4/4 seeds）"),
               ("打乱对照（溯源与旁信息置换）", "Δ ≈ 0", "+0.0007 / +0.0011 AUC，增益不是来自列数"),
               ("增益落在哪里", "稀疏 miRNA", "支持度 1–7 的正例损失下降 8–15%，16+ 持平")],
        source="数值来自 REPORT_WARM_PHL.md §0、§4.1、§4.3；基线 = ctx+diff × LR（K=10）；对 PLMF-MDA 官方复现领先 +0.033~0.035 AUC / +0.040~0.042 AUPR。",
        diag_title="诊断：warm 残差集中在稀疏 miRNA，PHL 的增益也落在那里",
        diag_sub="seed 42，论文 5 折。损失 = 排在每个测试正例前面的负例比例（均值 = 1 − AUC）；按 miRNA 在训练支持图中的度数分桶。",
        cats=["0", "1", "2–3", "4–7", "8–15", "16+"],
        a_series=("正例占比 %", "损失占比 %"),
        b_series=("基线 ctx+diff × LR", "PHL"),
        a_titles=("MDR · 正例占比 vs 损失占比", "MDS · 正例占比 vs 损失占比"),
        b_titles=("MDR · 正例侧单例均损：基线 → PHL", "MDS · 正例侧单例均损：基线 → PHL"),
        a_notes=("度 ≤3：正例 40%，损失 55%", "度 ≤3：正例 40%，损失 57%"),
        b_notes=("度 1–7 相对下降 13–15%；16+ 持平", "度 1–7 相对下降 8–14%；16+ 持平"),
        xlab="miRNA 在训练支持图中的度数",
        diag_source="数值来自 REPORT_WARM_PHL.md §1.1、§4.3（error_analysis.py）。",
    ),
    "en": dict(
        title="PHL: Provenance-Hypergraph Heuristic Learner — single-model architecture",
        subtitle="Curated provenance hyperedges and miRNA side-information relations are stacked on the PLMF-MDA support graph; "
                 "128 deterministic heuristics feed one logistic-regression readout (~130 parameters). Single model, linear, leak-free.",
        cols=["① Data sources", "② Multi-relational hypergraph (support only)",
              "③ Heuristic features (128 cols)", "④ Linear readout"],
        rows=[
            dict(src=("PLMF-MDA training positives", "miRNA–drug association matrix (MDR / MDS), training-fold positives only"),
                 g_title="Support graph + kNN content graph",
                 g_cap=["Bipartite support graph: training positives only", "kNN content graph: miRNA / drug similarity",
                        "solid = support edge; dashed = kNN edge"],
                 f_title="ctx (33) + diff (52)",
                 f_body=["degree, common neighbours, path3/5, SVD16, attribute support",
                         "PPR / Katz / path7 with row / column normalisation", "(original iNcRD-HG feature bank)"]),
            dict(src=("iNcRD-HG provenance records", "PMID of every association plus the cell-line / phenotype / disease context it was reported in"),
                 g_title="Provenance hypergraph",
                 g_cap=["Study hyperedge: all (m, d) reported by one PMID", "Context hyperedge: cell-line / phenotype / disease tokens",
                        "Co-study graph S = M·Mᵀ"],
                 f_title="prov (24)",
                 f_body=["PMID bridging |P_m ∩ P_d|, size-weighted, Jaccard (6)", "co-study transfer (S_r A), (A S_d) (6)",
                         "context-token overlap (4) · record statistics (8)"]),
            dict(src=("External miRNA knowledge (label-free)", "HMDD v4 disease spectrum · name / seed family · degree in the other phenotype matrix"),
                 g_title="miRNA side-information relations",
                 g_cap=["miRNA → family / disease / cross-phenotype degree",
                        "label-free; adds information for zero- and low-degree miRNAs"],
                 f_title="mprior (19)",
                 f_body=["name family (5) · seed 7-mer family (5)", "HMDD disease bridge Σ H[m,dz]·(HᵀA)[dz,d] (7)",
                         "degree in the other phenotype, xdeg (2)"]),
        ],
        g1_labels=("miRNA", "drug"),
        g2_labels=("PMID study hyperedge", "bridge m₁–d₂"),
        g3_labels=("miRNA", ["name / seed family", "HMDD diseases", "other-phenotype degree"]),
        note=[("Leak-free cross-fitting", dict(bold=True, size=9.5)),
              "Training rows: K=10 inner folds; a query pair never enters its own support",
              "Test rows: all training positives as support",
              "Every heuristic reads support positives only"],
        strip_label="Concatenate 128 columns",
        strip=["ctx+diff 85", "prov 24", "mprior 19"],
        lr=[("z = signed log1p + standardise", dict(size=9)),
            ("score(m, d) = w·z + b", dict(size=10.5, bold=True)),
            ("logistic regression, C = 0.3, ~130 parameters", dict(size=8.5, color=INK2))],
        out="score(m, d)",
        tiles=[("MDR · AUC / AUPR (paper 5-fold, 4 seeds)", "0.9591 / 0.9565", "+0.0062 / +0.0071 vs baseline 0.9529 / 0.9494 (4/4 seeds)"),
               ("MDS · AUC / AUPR (paper 5-fold, 4 seeds)", "0.9607 / 0.9587", "+0.0047 / +0.0051 vs baseline 0.9560 / 0.9536 (4/4 seeds)"),
               ("Shuffle control (provenance + side info permuted)", "Δ ≈ 0", "+0.0007 / +0.0011 AUC: the gain is not the extra columns"),
               ("Where the gain lands", "sparse miRNAs", "loss on support-degree 1–7 positives drops 8–15%; 16+ flat")],
        source="Numbers from REPORT_WARM_PHL.md §0, §4.1, §4.3; baseline = ctx+diff × LR (K=10); +0.033–0.035 AUC / +0.040–0.042 AUPR over the PLMF-MDA official reproduction.",
        diag_title="Diagnosis: the warm residual sits on sparse miRNAs, and so does the PHL gain",
        diag_sub="seed 42, paper 5-fold. Loss = fraction of negatives ranked above each test positive (mean = 1 − AUC); bucketed by miRNA degree in the training support graph.",
        cats=["0", "1", "2–3", "4–7", "8–15", "16+"],
        a_series=("share of positives %", "share of loss %"),
        b_series=("baseline ctx+diff × LR", "PHL"),
        a_titles=("MDR · share of positives vs share of loss", "MDS · share of positives vs share of loss"),
        b_titles=("MDR · mean loss per positive, baseline vs PHL", "MDS · mean loss per positive, baseline vs PHL"),
        a_notes=("degree ≤3: 40% of positives, 55% of loss", "degree ≤3: 40% of positives, 57% of loss"),
        b_notes=("degree 1–7: 13–15% relative drop; 16+ flat", "degree 1–7: 8–14% relative drop; 16+ flat"),
        xlab="miRNA degree in the training support graph",
        diag_source="Numbers from REPORT_WARM_PHL.md §1.1, §4.3 (error_analysis.py).",
    ),
}

# ------------------------------------------------------------------- data
DEG_POS = {"MDR": [10.0, 10.4, 19.9, 29.2, 26.8, 3.7], "MDS": [10.7, 11.5, 17.7, 28.9, 28.1, 3.1]}
DEG_LOSS = {"MDR": [13.4, 19.6, 22.5, 26.6, 15.8, 2.1], "MDS": [12.3, 22.3, 22.3, 25.6, 16.1, 1.4]}
GAIN_BASE = {"MDR": [0.054, 0.094, 0.055, 0.045, 0.030, 0.029], "MDS": [0.043, 0.088, 0.058, 0.041, 0.027, 0.022]}
GAIN_PHL = {"MDR": [0.050, 0.081, 0.047, 0.039, 0.027, 0.031], "MDS": [0.038, 0.076, 0.053, 0.036, 0.025, 0.023]}

HUES = [(BLUE, BLUE_T), (ORANGE, ORANGE_T), (AQUA, AQUA_T)]

# geometry (inches)
C1, W1 = 0.45, 2.35
C2, W2 = 3.25, 3.55
C3, W3 = 7.25, 2.45
C4, W4 = 10.15, 2.75
ROW_Y = [1.68, 3.13, 4.58]
ROW_H = 1.28


# --------------------------------------------------------------- schematics
def draw_g1(slide, px, py, hue, labels):
    mx, dx = px + 0.38, px + 1.42
    ys = [py + 0.56, py + 0.84, py + 1.12]
    add_text(slide, mx - 0.35, py + 0.30, 0.7, 0.2, [labels[0]], size=7.5, color=MUTED, align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))
    add_text(slide, dx - 0.35, py + 0.30, 0.7, 0.2, [labels[1]], size=7.5, color=MUTED, align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))
    r = 0.065
    for (i, j) in [(0, 0), (0, 1), (1, 1), (2, 2)]:
        add_line(slide, mx + r, ys[i], dx - r, ys[j], color=hue, width=1.25)
    # kNN content edges (dashed)
    add_line(slide, mx, ys[1] + r, mx, ys[2] - r, color=INK2, width=1.0, dash="sysDash")
    add_line(slide, dx, ys[0] + r, dx, ys[1] - r, color=INK2, width=1.0, dash="sysDash")
    for y in ys:
        node_circle(slide, mx, y, r, fill=hue)
        node_square(slide, dx, y, r, fill=SURF, line=hue)


def draw_g2(slide, px, py, hue, labels):
    ox, oy, ow, oh = px + 0.12, py + 0.36, 1.78, 0.86
    add_box(slide, ox, oy, ow, oh, fill=hue, alpha=16, shape=MSO_SHAPE.OVAL)
    add_text(slide, ox, oy + 0.02, ow, 0.18, [labels[0]], size=7.5, color=INK2, align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))
    mx, dx = px + 0.50, px + 1.52
    y1, y2 = py + 0.66, py + 1.04
    r = 0.065
    add_line(slide, mx + r, y1, dx - r, y1, color=INK2, width=1.25)
    add_line(slide, mx + r, y2, dx - r, y2, color=INK2, width=1.25)
    add_line(slide, mx + r, y1 + 0.02, dx - r, y2 - 0.02, color=hue, width=1.5, dash="sysDash")
    add_text(slide, px + 0.40, py + 0.87, 0.62, 0.13, [labels[1]], size=6.5, color=INK2, align=PP_ALIGN.RIGHT,
             margins=(0, 0, 0, 0), wrap=False)
    node_circle(slide, mx, y1, r, fill=hue)
    node_circle(slide, mx, y2, r, fill=hue)
    node_square(slide, dx, y1, r, fill=SURF, line=hue)
    node_square(slide, dx, y2, r, fill=SURF, line=hue)
    add_text(slide, mx - 0.4, y1 - 0.09, 0.3, 0.16, ["m₁"], size=7, color=INK2, align=PP_ALIGN.RIGHT, margins=(0, 0, 0, 0))
    add_text(slide, mx - 0.4, y2 - 0.09, 0.3, 0.16, ["m₂"], size=7, color=INK2, align=PP_ALIGN.RIGHT, margins=(0, 0, 0, 0))
    add_text(slide, dx + 0.1, y1 - 0.09, 0.3, 0.16, ["d₁"], size=7, color=INK2, margins=(0, 0, 0, 0))
    add_text(slide, dx + 0.1, y2 - 0.09, 0.3, 0.16, ["d₂"], size=7, color=INK2, margins=(0, 0, 0, 0))


def draw_g3(slide, px, py, hue, labels):
    cx, cy, r = px + 0.32, py + 0.80, 0.085
    sx, sw, sh = px + 0.72, 1.16, 0.22
    sys_ = [py + 0.41, py + 0.69, py + 0.97]
    for y, lab in zip(sys_, labels[1]):
        add_line(slide, cx + r, cy, sx, y + sh / 2, color=hue, width=1.25)
        b = add_box(slide, sx, y, sw, sh, fill=AQUA_T if hue == AQUA else SURF, line=None)
        fill_text(b.text_frame, [lab], size=7, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                  margins=(0.02, 0, 0.02, 0))
    node_circle(slide, cx, cy, r, fill=hue)
    add_text(slide, cx - 0.4, cy + r + 0.01, 0.8, 0.18, [labels[0]], size=7.5, color=MUTED, align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))


# ------------------------------------------------------------- architecture
def build_architecture(prs, lang):
    T = S[lang]
    s = new_slide(prs)
    add_text(s, 0.45, 0.25, 12.4, 0.48, [T["title"]], size=20, bold=True)
    add_text(s, 0.45, 0.74, 12.4, 0.5, [T["subtitle"]], size=10, color=INK2)

    for x, w, txt in [(C1, W1, T["cols"][0]), (C2, W2, T["cols"][1]), (C3, W3, T["cols"][2]), (C4, W4, T["cols"][3])]:
        add_text(s, x, 1.30, w, 0.3, [txt], size=10, bold=True, color=INK2, margins=(0.02, 0, 0.02, 0))
        add_line(s, x, 1.62, x + w, 1.62, color=BASE, width=0.75)

    strip_yc = ROW_Y[1] + ROW_H / 2
    for i, row in enumerate(T["rows"]):
        hue, tint = HUES[i]
        y = ROW_Y[i]
        yc = y + ROW_H / 2
        # col 1: data source (tinted box)
        b = add_box(s, C1, y, W1, ROW_H, fill=tint)
        fill_text(b.text_frame, [(row["src"][0], dict(bold=True, size=9.5)), (row["src"][1], dict(size=8.5, color=INK2))],
                  anchor=MSO_ANCHOR.MIDDLE, margins=(0.12, 0.06, 0.1, 0.06), space_after=3)
        # col 2: structure panel (white, colored hairline) + schematic + caption
        add_box(s, C2, y, W2, ROW_H, fill=SURF, line=hue, line_w=1.0)
        node_circle(s, C2 + 0.17, y + 0.17, 0.05, fill=hue)
        add_text(s, C2 + 0.26, y + 0.04, 2.6, 0.26, [row["g_title"]], size=9.5, bold=True, margins=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
        [draw_g1, draw_g2, draw_g3][i](s, C2, y, hue, T[["g1_labels", "g2_labels", "g3_labels"][i]])
        add_text(s, C2 + 1.95, y + 0.34, W2 - 2.02, ROW_H - 0.38, row["g_cap"], size=7.5, color=INK2,
                 margins=(0.02, 0, 0.02, 0), space_after=2)
        # col 3: feature group (white, left accent bar, hairline)
        add_box(s, C3, y, W3, ROW_H, fill=SURF, line=GRID, line_w=0.75, radius=0.06)
        add_box(s, C3, y + 0.10, 0.07, ROW_H - 0.20, fill=hue, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, C3 + 0.16, y + 0.05, W3 - 0.22, ROW_H - 0.1,
                 [(row["f_title"], dict(bold=True, size=9.5))] + [(t, dict(size=7.5, color=INK2)) for t in row["f_body"]],
                 margins=(0, 0.02, 0, 0.02), space_after=2, anchor=MSO_ANCHOR.MIDDLE)
        # arrows col1 -> col2 -> col3
        add_line(s, C1 + W1, yc, C2, yc, head=True)
        add_line(s, C2 + W2, yc, C3, yc, head=True)
        # col3 -> strip (row 2 straight, others elbow)
        kind = MSO_CONNECTOR.STRAIGHT if i == 1 else MSO_CONNECTOR.ELBOW
        add_line(s, C3 + W3, yc, C4, strip_yc, head=True, kind=kind)

    # col 4: note (row 1), strip (row 2), LR + score (row 3)
    b = add_box(s, C4, ROW_Y[0], W4, ROW_H, fill=TILE, line=GRID, line_w=0.75, radius=0.06)
    fill_text(b.text_frame, [T["note"][0]] + [(t, dict(size=7.5, color=INK2)) for t in T["note"][1:]],
              margins=(0.1, 0.05, 0.1, 0.05), space_after=2, anchor=MSO_ANCHOR.MIDDLE)
    sy, sh = strip_yc - 0.17, 0.34
    add_text(s, C4, sy - 0.24, W4, 0.22, [T["strip_label"]], size=8, color=MUTED, margins=(0.02, 0, 0, 0), align=PP_ALIGN.CENTER)
    widths = [85, 24, 19]
    x = C4
    gap = 0.03
    usable = W4 - 2 * gap
    for j, (wv, (hue, tint)) in enumerate(zip(widths, HUES)):
        w = usable * wv / 128
        add_box(s, x, sy, w, sh, fill=hue, shape=MSO_SHAPE.RECTANGLE)
        x += w + gap
    # legend row under the strip (identity never rides on colour alone)
    lx = C4 + 0.02
    for j, (hue, tint) in enumerate(HUES):
        node_circle(s, lx + 0.05, sy + sh + 0.13, 0.04, fill=hue)
        add_text(s, lx + 0.12, sy + sh + 0.04, 0.8, 0.18, [T["strip"][j]], size=7.5, color=INK2, margins=(0, 0, 0, 0))
        lx += 0.93
    ly, lh = 4.42, 0.88
    add_line(s, C4 + W4 / 2, sy + sh + 0.24, C4 + W4 / 2, ly, head=True)
    b = add_box(s, C4, ly, W4, lh, fill=SURF, line=INK2, line_w=1.0, radius=0.08)
    fill_text(b.text_frame, T["lr"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0.06, 0.04, 0.06, 0.04), space_after=2)
    oy, oh = 5.50, 0.36
    add_line(s, C4 + W4 / 2, ly + lh, C4 + W4 / 2, oy, head=True)
    b = add_box(s, C4 + 0.55, oy, W4 - 1.1, oh, fill=INK, radius=0.3)
    fill_text(b.text_frame, [T["out"]], size=11, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))

    # footer: stat tiles
    ty, th, tw, tg = 6.18, 0.92, 3.02, 0.11
    for k, (lab, val, delta) in enumerate(T["tiles"]):
        x = 0.45 + k * (tw + tg)
        b = add_box(s, x, ty, tw, th, fill=TILE, line=GRID, line_w=0.75, radius=0.06)
        fill_text(b.text_frame, [(lab, dict(size=7.5, color=MUTED)), (val, dict(size=14, bold=True)),
                                 (delta, dict(size=7.5, color=GOOD if k < 2 else INK2))],
                  margins=(0.12, 0.05, 0.1, 0.04), space_after=1, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 0.45, 7.14, 12.4, 0.25, [T["source"]], size=7, color=MUTED, margins=(0.02, 0, 0, 0))
    return s


# --------------------------------------------------------------- diagnostics
def _patch_chart_fonts(chart, name=FONT):
    for tag in ("a:defRPr", "a:rPr", "a:endParaRPr"):
        for el in chart._chartSpace.iter(qn(tag)):
            latin = el.find(qn("a:latin"))
            if latin is None:
                latin = etree.SubElement(el, qn("a:latin"))
            latin.set("typeface", name)
            _set_ea(el, name)


def _chart_no_border(chart):
    cs = chart._chartSpace
    spPr = etree.Element(qn("c:spPr"))
    etree.SubElement(spPr, qn("a:noFill"))
    ln = etree.SubElement(spPr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    cs.find(qn("c:chart")).addnext(spPr)


def add_bar_chart(slide, x, y, w, h, cats, series, colors, ymax, numfmt, xlab=None, major=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(8)
    ch.font.name = FONT
    ch.font.color.rgb = rgb(INK2)
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(8)
    plot = ch.plots[0]
    plot.gap_width = 240
    plot.overlap = -15
    plot.vary_by_categories = False
    for sr, col in zip(plot.series, colors):
        sr.format.fill.solid()
        sr.format.fill.fore_color.rgb = rgb(col)
        sr.format.line.fill.background()
    va = ch.value_axis
    va.minimum_scale = 0
    va.maximum_scale = ymax
    if major is not None:
        va.major_unit = major
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(GRID)
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.font.size = Pt(8)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.tick_labels.number_format = numfmt
    va.tick_labels.number_format_is_linked = False
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(BASE)
    ca.format.line.width = Pt(0.75)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(8)
    ca.tick_labels.font.color.rgb = rgb(INK2)
    if xlab:
        ca.has_title = True
        ca.axis_title.text_frame.text = xlab
        r = ca.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(8)
        r.font.bold = False
        r.font.color.rgb = rgb(MUTED)
    _chart_no_border(ch)
    _patch_chart_fonts(ch)
    return ch


def build_diagnostics(prs, lang):
    T = S[lang]
    s = new_slide(prs)
    add_text(s, 0.45, 0.25, 12.4, 0.48, [T["diag_title"]], size=20, bold=True)
    add_text(s, 0.45, 0.74, 12.4, 0.45, [T["diag_sub"]], size=10, color=INK2)
    cells = [(0.45, 1.35), (6.78, 1.35), (0.45, 4.25), (6.78, 4.25)]
    cw, chh = 6.1, 2.55
    specs = []
    for r_, ds in enumerate(["MDR", "MDS"]):
        specs.append(dict(title=T["a_titles"][r_], note=T["a_notes"][r_],
                          series=[(T["a_series"][0], DEG_POS[ds]), (T["a_series"][1], DEG_LOSS[ds])],
                          colors=[BASE, BLUE], ymax=30, fmt='0"%"', cell=cells[2 * r_], major=10))
        specs.append(dict(title=T["b_titles"][r_], note=T["b_notes"][r_],
                          series=[(T["b_series"][0], GAIN_BASE[ds]), (T["b_series"][1], GAIN_PHL[ds])],
                          colors=[BASE, BLUE], ymax=0.10, fmt="0.00", cell=cells[2 * r_ + 1], major=0.02))
    for k, sp in enumerate(specs):
        x, y = sp["cell"]
        add_text(s, x, y, cw * 0.6, 0.3, [sp["title"]], size=9.5, bold=True, margins=(0.02, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + cw * 0.6, y, cw * 0.4, 0.3, [sp["note"]], size=8, color=INK2, align=PP_ALIGN.RIGHT,
                 margins=(0, 0, 0.02, 0), anchor=MSO_ANCHOR.MIDDLE)
        add_bar_chart(s, x, y + 0.3, cw, chh, T["cats"], sp["series"], sp["colors"], sp["ymax"], sp["fmt"],
                      xlab=T["xlab"] if k >= 2 else None, major=sp["major"])
    add_text(s, 0.45, 7.14, 12.4, 0.25, [T["diag_source"]], size=7, color=MUTED, margins=(0.02, 0, 0, 0))
    return s


def main(out):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    build_architecture(prs, "cn")
    build_diagnostics(prs, "cn")
    build_architecture(prs, "en")
    build_diagnostics(prs, "en")
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    import sys
    main(sys.argv[1] if len(sys.argv) > 1 else "PHL_architecture.pptx")
