# -*- coding: utf-8 -*-
"""Build the DeepLDA reproduction Excel(s).

1) deeplda_official_20260920/DeepLDA复现汇总.xlsx  (standalone workbook)
2) append a sheet + a ledger row to
   improve/results/external_specialized_20260916/外部专用模型对比_论文格式.xlsx
"""
import csv
import glob
import os
import statistics as st

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
MAIN_XLSX = "/data1/licheng/incrd-hg/iNcRD-HG/improve/results/external_specialized_20260916/外部专用模型对比_论文格式.xlsx"

PAPER = {
    ("1", 0): (0.9536, 0.8511, 0.8612, 0.7562),
    ("1", 1): (0.9583, 0.8601, 0.8625, 0.7628),
    ("2", 0): (0.9728, 0.8572, 0.8876, 0.7792),
    ("2", 1): (0.9819, 0.8687, 0.9163, 0.7885),
}
PAPER_AVG = {("1", 1): {"gat": 0.8289, "gcn": 0.7808},
             ("1", 0): {"gat": 0.7820, "gcn": 0.7541},
             ("2", 1): {"gat": 0.7949, "gcn": 0.7769},
             ("2", 0): {"gat": 0.7740, "gcn": 0.7466}}
DSNAME = {"1": "NoncoRNA (Dataset1, 3601x71/3802)", "2": "ncDR (Dataset2, 162x31/184)"}

HEAD = Font(bold=True, color="FFFFFF", size=10)
HEADF = PatternFill("solid", fgColor="305496")
OURF = PatternFill("solid", fgColor="FFF2CC")
POS = PatternFill("solid", fgColor="FFC7CE")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)


def load_groups():
    groups = {}
    detail = []
    for fp in glob.glob(os.path.join(HERE, "results_*.csv")):
        for r in csv.DictReader(open(fp)):
            bal = int(r.get("balanced", 0))
            key = (str(r["dataset"]), bal, r["variant"])
            vals = {k: float(r[k]) for k in ("auc", "aupr", "f1", "mcc")}
            groups.setdefault(key, []).append(vals)
            detail.append((str(r["dataset"]), bal, r["variant"], int(r["seed"]), vals))
    return groups, detail


def ms(x):
    return f"{st.mean(x):.4f}±{st.pstdev(x):.4f}"


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD
        cell.fill = HEADF
        cell.alignment = CEN
        cell.border = BORD


def build_standalone(groups, detail):
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    notes = [
        ["项", "说明"],
        ["目标", "复现 DeepLDA (Gao & Shang, Front. Microbiol. 2023, doi:10.3389/fmicb.2023.1147778) 基线"],
        ["方法", "GCN(多属性特征提取) + GAT(注意力聚合) -> 内积解码"],
        ["官方代码", "github.com/meihonggao/DeepLDA (official_codes/, 含 1 行 NameError 修复 super(MVG_LGA)->super(Deep_LDA))"],
        ["官方数据", "Datasets/Dataset1=NoncoRNA 3601x71/3802; Dataset2=ncDR 162x31/184"],
        ["官方协议", "打乱行 -> 用整张关联矩阵算高斯相似度当节点特征(标签泄漏) -> 前90%行训练/后10%行测试, 全部元素打分"],
        ["balanced", "官方代码未实现; 按论文文字补: 正例 + 等量随机负例, 训练/测试各自抽样"],
        ["运行设置", "5 seeds [72,666,0,1,2], epochs=1050, lr=1e-3, wd=5e-4, nhid=50, nclass=200, nheads=8, dropout=0.0"],
        ["结论1", "ncDR 部分复现(同量级, 方差大); NoncoRNA 未复现(unbal 退化为全负 MCC=0, bal AUC 0.60)"],
        ["结论2", "官方泄漏协议下 GCN/GAT 单分支普遍 >= DeepLDA; 论文 'DeepLDA>GAT>GCN' 未复现"],
        ["结论3", "官方 F1/MCC 用 0.5 阈值, 极不平衡时退化; AUC 相对稳定"],
        ["产物", "results_*.csv / REPORT.md / SUMMARY_deeplda_official.md"],
    ]
    for r in notes:
        ws.append(r)
    style_header(ws, 1, 2)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110

    ws2 = wb.create_sheet("官方复现-论文口径")
    hdr = ["数据集", "设置", "变体", "AUC", "AUPR", "F1", "MCC", "论文DeepLDA(AUC/AUPR/F1/MCC)", "论文GCN/GAT四指标均值", "备注"]
    ws2.append(hdr)
    style_header(ws2, 1, len(hdr))
    for ds in ("1", "2"):
        for bal in (0, 1):
            setting = "balanced" if bal else "unbalanced"
            pa, pp, pf, pm = PAPER[(ds, bal)]
            for v in ("deeplda", "gcn", "gat"):
                g = groups.get((ds, bal, v), [])
                row = [DSNAME[ds], setting, v]
                if g:
                    for k in ("auc", "aupr", "f1", "mcc"):
                        row.append(ms([x[k] for x in g]))
                else:
                    row += ["-", "-", "-", "-"]
                if v == "deeplda":
                    row += [f"{pa:.4f}/{pp:.4f}/{pf:.4f}/{pm:.4f}", "-",
                            "官方代码; 未复现" if ds == "1" else "官方代码; 同量级"]
                else:
                    pavg = PAPER_AVG.get((ds, bal), {}).get(v)
                    row += ["-", f"{pavg:.4f}" if pavg else "-", "同协议单分支对照"]
                ws2.append(row)
    for r in range(2, ws2.max_row + 1):
        for c in range(1, len(hdr) + 1):
            ws2.cell(row=r, column=c).border = BORD
            ws2.cell(row=r, column=c).alignment = LEFT if c in (1, 2, 3, 8, 9, 10) else CEN
        if ws2.cell(row=r, column=3).value == "deeplda":
            for c in range(1, len(hdr) + 1):
                ws2.cell(row=r, column=c).fill = OURF
    widths = [30, 11, 9, 13, 13, 13, 13, 30, 18, 26]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("逐种子明细")
    ws3.append(["数据集", "设置", "变体", "seed", "AUC", "AUPR", "F1", "MCC"])
    style_header(ws3, 1, 8)
    for ds, bal, v, seed, vals in sorted(detail, key=lambda x: (x[0], x[1], x[2], x[3])):
        ws3.append([ds, "balanced" if bal else "unbalanced", v, seed,
                    vals["auc"], vals["aupr"], vals["f1"], vals["mcc"]])
    for r in range(2, ws3.max_row + 1):
        for c in range(1, 9):
            ws3.cell(row=r, column=c).border = BORD
            ws3.cell(row=r, column=c).alignment = CEN
            if c >= 5:
                ws3.cell(row=r, column=c).number_format = "0.0000"
    for i, w in enumerate([10, 11, 9, 7, 10, 10, 10, 10], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws4 = wb.create_sheet("同协议-ncDR701x101")
    ws4.append(["方法", "AUC", "AUPR", "ACC", "F1", "MCC", "协议/备注"])
    style_header(ws4, 1, 7)
    ws4.append(["DeepLDA (本地移植)", "0.6788±0.0529", "0.7457±0.0315", "0.6851±0.0952",
                "0.7124±0.0380", "0.3714±0.1906",
                "common.py 同协议: 701 miRNA x 101 drug/3163, RandomLinkSplit(test=0.2), 等量均匀负采样, 5 seeds; repro_deeplda_ncdr.py"])
    for r in range(2, ws4.max_row + 1):
        for c in range(1, 8):
            ws4.cell(row=r, column=c).border = BORD
            ws4.cell(row=r, column=c).alignment = LEFT if c in (1, 7) else CEN
    for i, w in enumerate([22, 14, 14, 14, 14, 14, 80], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    out = os.path.join(HERE, "DeepLDA复现汇总.xlsx")
    wb.save(out)
    print("saved", out)
    return out


def update_main(groups):
    wb = load_workbook(MAIN_XLSX)
    sheet = "DeepLDA官方复现"
    if sheet in wb.sheetnames:
        del wb[sheet]
    ws = wb.create_sheet(sheet)
    hdr = ["数据集", "设置", "变体", "AUC", "AUPR", "F1", "MCC", "论文DeepLDA(AUC/AUPR/F1/MCC)", "备注"]
    ws.append(["DeepLDA (Front. Microbiol. 2023) 官方代码复现  |  GCN + GAT  |  5 seeds"])
    ws.append(hdr)
    style_header(ws, 2, len(hdr))
    for ds in ("1", "2"):
        for bal in (0, 1):
            setting = "balanced" if bal else "unbalanced"
            pa, pp, pf, pm = PAPER[(ds, bal)]
            for v in ("deeplda", "gcn", "gat"):
                g = groups.get((ds, bal, v), [])
                row = [DSNAME[ds], setting, v]
                if g:
                    for k in ("auc", "aupr", "f1", "mcc"):
                        row.append(ms([x[k] for x in g]))
                else:
                    row += ["-", "-", "-", "-"]
                if v == "deeplda":
                    row += [f"{pa:.4f}/{pp:.4f}/{pf:.4f}/{pm:.4f}",
                            "官方代码; NoncoRNA未复现, ncDR同量级" if ds == "1" else "官方代码; 同量级, 方差大"]
                else:
                    row += ["-", "同协议单分支对照(泄漏协议下>=DeepLDA)"]
                ws.append(row)
    for r in range(3, ws.max_row + 1):
        for c in range(1, len(hdr) + 1):
            ws.cell(row=r, column=c).border = BORD
            ws.cell(row=r, column=c).alignment = LEFT if c in (1, 2, 3, 8, 9) else CEN
        if ws.cell(row=r, column=3).value == "deeplda":
            for c in range(1, len(hdr) + 1):
                ws.cell(row=r, column=c).fill = OURF
    for i, w in enumerate([30, 11, 9, 13, 13, 13, 13, 30, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    if "同协议复现台账" in wb.sheetnames:
        lt = wb["同协议复现台账"]
        if not any(lt.cell(row=r, column=1).value == "DeepLDA" for r in range(1, lt.max_row + 1)):
            lt.append(["DeepLDA", "Front Microbiol 2023", "lncRNA-drug 抗性",
                       "github.com/meihonggao/DeepLDA", "torch1.13+pyg (venv-dgl)",
                       "已完成(官方代码)",
                       "AUC 0.6788±0.0529 / AUPR 0.7457±0.0315",
                       "官方代码修 1 处 NameError; 同协议本地移植 701x101; 论文口径见 DeepLDA官方复现 sheet"])
    wb.save(MAIN_XLSX)
    print("updated", MAIN_XLSX)


if __name__ == "__main__":
    groups, detail = load_groups()
    build_standalone(groups, detail)
    update_main(groups)
