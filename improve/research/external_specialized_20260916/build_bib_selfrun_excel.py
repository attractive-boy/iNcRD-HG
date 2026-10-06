"""Focused workbook: only B-journal / BIB baselines, only self-run values."""
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = ("/data1/licheng/incrd-hg/iNcRD-HG/improve/results/"
       "external_specialized_20260916/B刊-BIB基线_自跑.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

# dataset, method, venue, tier, task, AUC, AUPR, ACC, Prec, Rec, Spec, F1, MCC, how
ROWS = [
    ("MPHGNN 独立测试集", "MPHGNN", "Briefings in Bioinformatics", "CCF-B/中科院1区", "miRNA-drug 抗性",
     0.9755, 0.9415, 0.9246, 0.8883, 0.8850, 0.9443, 0.8866, 0.8301, "官方代码, 5折 seed42"),
    ("MPHGNN 独立测试集", "SubKNet", "IEEE JBHI", "CCF-B/中科院1区", "miRNA-drug 抗性",
     0.9493, 0.9070, None, None, None, None, None, None, "官方模型复现, 5折 seed42; 阈值指标接口不稳定"),
    ("ncDR (GAM-MDR 701x101)", "GAM-MDR", "Briefings in Functional Genomics", "中科院2区", "miRNA-drug 抗性",
     0.9179, 0.8977, 0.8540, 0.8299, 0.8908, 0.8171, 0.8590, 0.7104, "官方代码移植, 5 seeds"),
    ("ncDR (GAM-MDR 701x101)", "NDSGCL", "IEEE/ACM TCBB", "CCF-B/中科院2区", "ncRNA-drug 抗性",
     0.9004, 0.9070, 0.8282, 0.8254, 0.8335, 0.8228, 0.8291, 0.6569, "重实现, 5 seeds"),
    ("ncDR (GAM-MDR 701x101)", "SGCLDGA", "Briefings in Bioinformatics", "CCF-B/中科院1区", "drug-gene",
     0.8878, 0.9089, 0.8378, 0.8366, 0.8408, 0.8348, 0.8384, 0.6761, "官方代码, 5 seeds"),
    ("ncDR (GAM-MDR 701x101)", "GCMDR", "Bioinformatics", "CCF-B", "miRNA-drug 抗性",
     0.7875, 0.6887, 0.7763, 0.7649, 0.8079, 0.7446, 0.7838, 0.5567, "PyG 重实现, 5 seeds"),
    ("ncDR (GAM-MDR 701x101)", "DLP", "IEEE/ACM TCBB", "CCF-B/中科院2区", "drug-miRNA",
     0.7622, 0.8219, 0.7419, None, None, None, 0.7414, 0.4858, "MATLAB→Python 移植, 5 seeds"),
    ("ncDR (GAM-MDR 701x101)", "AGCLNDA", "IEEE JBHI", "CCF-B/中科院1区", "ncRNA-drug 抗性",
     0.6892, 0.7485, None, None, None, None, None, None, "官方代码, 5 seeds; 阈值指标退化"),
    ("ncDR (GAM-MDR 701x101)", "JMSSMMA", "Briefings in Bioinformatics", "CCF-B/中科院1区", "small molecule-miRNA",
     0.6145, 0.6305, None, None, None, None, None, None, "官方代码, 5 seeds; 阈值指标退化"),
    ("ncDR (GAM-MDR 701x101)", "iNcRD-HG (本文)", "Bioinformatics", "CCF-B", "ncRNA-drug 抗性",
     0.8322, 0.8135, 0.7718, 0.7404, 0.8427, 0.7009, 0.7869, 0.5515, "图特征 + 冻结 TabICLv2, 5 seeds"),
    ("GCFMCL D2R", "iNcRD-HG (本文)", "Bioinformatics", "CCF-B", "miRNA-drug 敏感性",
     0.7482, 0.7492, 0.6759, 0.6341, 0.8472, 0.5046, 0.7236, 0.3767, "图特征 + 冻结 TabICLv2, 5折"),
    ("NoncoRNA+ncDR 41 lncRNA", "iNcRD-HG (本文)", "Bioinformatics", "CCF-B", "lncRNA-drug 抗性",
     0.7234, 0.7347, 0.6677, 0.6114, 0.9308, 0.4047, 0.7376, 0.3928, "图特征 + 冻结 TabICLv2, 5折"),
    ("NoncoRNA+ncDR 2693", "iNcRD-HG (本文)", "Bioinformatics", "CCF-B", "ncRNA-drug 抗性",
     0.7617, 0.7782, 0.6951, 0.6610, 0.8058, 0.5845, 0.7255, 0.4012, "图特征 + 冻结 TabICLv2, 5折"),
    ("RNAactDrug (HGNNLDA 数据)", "iNcRD-HG (本文)", "Bioinformatics", "CCF-B", "lncRNA-drug 敏感性",
     0.9016, 0.9144, 0.8323, 0.8332, 0.8310, 0.8336, 0.8321, 0.6647, "图特征 + 冻结 TabICLv2"),
]
COLS = ["数据集", "方法", "期刊", "分区", "任务", "AUC", "AUPR", "ACC",
        "Precision", "Recall", "Specificity", "F1", "MCC", "复现方式"]

DARK = PatternFill("solid", fgColor="305496")
OURS = PatternFill("solid", fgColor="FFF2CC")
HEAD = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "B-BIB基线(自跑)"
ws.append(COLS)
for row in ROWS:
    ws.append(list(row))
for c in range(1, len(COLS) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = DARK; cell.font = HEAD; cell.alignment = CENTER; cell.border = BORDER
for r in range(2, ws.max_row + 1):
    ours = "本文" in str(ws.cell(row=r, column=2).value)
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER
        cell.alignment = LEFT if c in (1, 2, 3, 4, 5, 14) else CENTER
        if c >= 6 and c <= 13 and isinstance(cell.value, (int, float)):
            cell.number_format = "0.0000"
        if ours:
            cell.fill = OURS
    ws.cell(row=r, column=2).font = BOLD
for i, w in enumerate([24, 20, 26, 16, 18, 9, 9, 9, 10, 9, 10, 9, 9, 30], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
wb.save(OUT)
print("wrote", OUT)
