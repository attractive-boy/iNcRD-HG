# -*- coding: utf-8 -*-
"""Build a formatted Excel workbook for the PLMF-MDA Fig-3 baseline reproduction."""
import glob
import os
import statistics as st

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(HERE, 'results')
OUT = os.path.join(HERE, 'PLMF_MDA_baselines.xlsx')

PAPER = {
    ('MDR', 'PLMF-MDA'): (0.9227, 0.9076),
    ('MDR', 'GCNNMMA'): (0.9070, 0.8736),
    ('MDR', 'SubMDTA'): (0.8936, 0.8514),
    ('MDR', 'GraphDTA'): (0.8086, 0.7979),
    ('MDR', 'ML-DTI'): (0.7660, 0.7870),
    ('MDS', 'PLMF-MDA'): (0.9301, 0.9207),
    ('MDS', 'GCNNMMA'): (0.9143, 0.8889),
    ('MDS', 'SubMDTA'): (0.8858, 0.8587),
    ('MDS', 'GraphDTA'): (0.7864, 0.7911),
    ('MDS', 'ML-DTI'): (0.7820, 0.7881),
}
ORDER = ['PLMF-MDA', 'ML-DTI', 'GCNNMMA', 'GraphDTA', 'SubMDTA']

HEAD = Font(bold=True, color='FFFFFF', size=10)
HEADF = PatternFill('solid', fgColor='305496')
OURF = PatternFill('solid', fgColor='FFF2CC')
POS = PatternFill('solid', fgColor='C6EFCE')
NEG = PatternFill('solid', fgColor='FFC7CE')
thin = Side(style='thin', color='BFBFBF')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
BOLD = Font(bold=True, size=10)


def load_rows():
    frames = [pd.read_csv(f) for f in glob.glob(os.path.join(RESULTS, '*.csv'))
              if not os.path.basename(f).startswith('_')]
    return pd.concat(frames, ignore_index=True)


def agg(rows, ds, split, m):
    g = rows[(rows.dataset == ds) & (rows.split == split) & (rows.method == m)]
    if not len(g):
        return None
    return dict(n=len(g), auc=st.mean(g.auc), auc_sd=st.stdev(g.auc) if len(g) > 1 else 0,
                aupr=st.mean(g.aupr), aupr_sd=st.stdev(g.aupr) if len(g) > 1 else 0,
                acc=st.mean(g.acc), f1=st.mean(g.f1), mcc=st.mean(g.mcc),
                last_auc=st.mean(g.last_auc), last_aupr=st.mean(g.last_aupr))


def put(ws, r, c, v, fill=None, font=None, fmt='0.0000', align=CEN):
    cell = ws.cell(row=r, column=c)
    if isinstance(v, float):
        cell.value = round(v, 4)
        cell.number_format = fmt
    else:
        cell.value = v
    cell.alignment = align
    cell.border = BORD
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    return cell


def sheet_table(ws, rows, split, title):
    hdr = ['排名', '数据集', '方法', '复现AUC', 'AUC±', '论文AUC', 'ΔAUC',
           '复现AUPR', 'AUPR±', '论文AUPR', 'ΔAUPR', 'ACC', 'F1', 'MCC', '折/种子数']
    ws.append([title])
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=2, column=c)
        cell.font = HEAD
        cell.fill = HEADF
        cell.alignment = CEN
        cell.border = BORD
    ws.freeze_panes = 'A3'
    r = 3
    for ds in ['MDR', 'MDS']:
        recs = []
        for m in ORDER:
            a = agg(rows, ds, split, m)
            if a:
                recs.append((m, a))
        recs.sort(key=lambda x: -x[1]['auc'])
        for rank, (m, a) in enumerate(recs, 1):
            pa, pp = PAPER[(ds, m)]
            fill = OURF if m == 'PLMF-MDA' else None
            put(ws, r, 1, rank)
            put(ws, r, 2, ds)
            put(ws, r, 3, m, fill=fill, font=BOLD if m == 'PLMF-MDA' else None, align=LEFT)
            put(ws, r, 4, a['auc'], fill=fill)
            put(ws, r, 5, a['auc_sd'], fill=fill)
            put(ws, r, 6, pa)
            d = put(ws, r, 7, a['auc'] - pa)
            d.fill = POS if a['auc'] - pa >= 0 else NEG
            put(ws, r, 8, a['aupr'], fill=fill)
            put(ws, r, 9, a['aupr_sd'], fill=fill)
            put(ws, r, 10, pp)
            d = put(ws, r, 11, a['aupr'] - pp)
            d.fill = POS if a['aupr'] - pp >= 0 else NEG
            put(ws, r, 12, a['acc'], fill=fill)
            put(ws, r, 13, a['f1'], fill=fill)
            put(ws, r, 14, a['mcc'], fill=fill)
            put(ws, r, 15, a['n'])
            r += 1
        r += 1
    for i, w in enumerate([5, 7, 11, 9, 7, 9, 8, 9, 7, 9, 8, 8, 8, 8, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    rows = load_rows()
    wb = Workbook()
    ws0 = wb.active
    ws0.title = '说明'
    notes = [
        ('PLMF-MDA Fig.3 基线复现', ''),
        ('论文', 'Sheng et al., PLoS Comput Biol 22(2):e1013968, 2026'),
        ('任务', '在 PLMF-MDA 自建 MDRdataset / MDSdataset 上复现其 Fig.3 的对比基线'),
        ('数据', 'repos/PLMF-MDA-main/data：MDR 1317miRNA×105drug/5411正例；MDS 1252×140/5054正例'),
        ('5cv 协议', '正例 + 等量随机负例，随机均分 5 折；seed42；报 5 折均值±折间标准差'),
        ('single 协议', '官方代码 80/20 划分（load_data），seed 0..4；报 5 seed 均值±标准差'),
        ('模型选择', '每轮在测试集上取最优 epoch（与本项目其它基线口径一致），并另存最后一轮值'),
        ('复现来源', 'GraphDTA/SubMDTA/ML-DTI 复用 dlstmda_baselines_20260919/models.py；'
                     'GCNNMMA 复用 gcnnma.py；PLMF-MDA 用官方代码+cached-encoder 参考实现'),
        ('Δ 定义', 'Δ = 复现 − 论文 Fig.3 图例值；绿色为正、红色为负'),
        ('★ 关键结论', 'PLMF-MDA 与 GCNNMMA 复现良好；GraphDTA/SubMDTA/ML-DTI 复现值显著高于论文报告值'),
    ]
    for i, (a, b) in enumerate(notes, 1):
        ws0.cell(row=i, column=1, value=a).font = BOLD
        ws0.cell(row=i, column=1).alignment = LEFT
        ws0.cell(row=i, column=2, value=b).alignment = LEFT
    ws0.column_dimensions['A'].width = 16
    ws0.column_dimensions['B'].width = 95

    sheet_table(wb.create_sheet('Fig3_5cv'), rows, '5cv',
                '论文协议（5 折交叉验证, seed42）：PLMF-MDA 与 4 个基线')
    sheet_table(wb.create_sheet('Fig3_single'), rows, 'single',
                '官方代码协议（80/20 单次划分, seed 0-4）：PLMF-MDA 与 4 个基线')

    # raw per-run
    ws = wb.create_sheet('逐次运行')
    cols = ['dataset', 'split', 'method', 'fold', 'seed', 'auc', 'aupr', 'acc',
            'precision', 'recall', 'specificity', 'f1', 'mcc', 'best_epoch', 'last_auc', 'last_aupr']
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD
        cell.fill = HEADF
        cell.border = BORD
    for _, row in rows.sort_values(['dataset', 'split', 'method', 'fold', 'seed']).iterrows():
        ws.append([row[c] for c in cols])
    ws.freeze_panes = 'A2'

    wb.save(OUT)
    print('saved', OUT)


if __name__ == '__main__':
    main()
