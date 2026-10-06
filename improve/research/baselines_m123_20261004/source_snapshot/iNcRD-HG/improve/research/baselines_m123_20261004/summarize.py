"""Build the requested delta tables, workbook, raw summaries and provenance."""
import csv
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from full_adapter import HERE, BENCH, SOURCE, load, sha
from audit import audit

ROOT=HERE.parents[3]
REFERENCE='incrd_original'


def write_csv(path, header, rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(header);w.writerows(rows)


def markdown(header, rows):
    def fmt(v):
        return f'{v:.4f}' if isinstance(v,(float,np.floating)) else str(v)
    return '\n'.join(['| '+' | '.join(header)+' |','| '+' | '.join(['---']*len(header))+' |']+
                     ['| '+' | '.join(map(fmt,r))+' |' for r in rows])


def main():
    raw=audit()
    for name in ('module_checks_initial.json','module_checks_trained.json'):
        checks=json.loads((HERE/name).read_text())
        assert set(checks)=={'MiDrug','MDR','MDS'}
    registry=json.loads((BENCH/'registry.json').read_text())
    registry.append(dict(id='incrd_m123',name='iNcRD-HG + M1 + M2 + M3',category='本项目改进模型',
        year='本次运行 2026-10-04',journal='—',doi='—',code=str(HERE/'full_adapter.py'),
        implementation='原始编码器及辅助图 + QMP + SP-RNCN + MC-EASE；MiDrug 按实际关系数推广 QMP；不含 M4',
        rationale='用户指定的完整模块组合'))
    names={r['id']:r['name'] for r in registry}; names[REFERENCE]='iNcRD-HG'
    order=[REFERENCE]+[r['id'] for r in registry if r['id']!=REFERENCE]
    datasets=('MiDrug','MDR','MDS')
    group={};summary=[]
    for m in order:
        for ds in datasets:
            rr=[r for r in raw if r['model']==m and r['dataset']==ds]
            group[m,ds]={k:float(np.mean([r[k] for r in rr])) for k in ('auc','aupr')}
            group[m,ds].update({k+'_sd':float(np.std([r[k] for r in rr],ddof=1)) for k in ('auc','aupr')})
            group[m,ds]['n']=len(rr)
    for m in order:
        for ds in datasets:
            g=group[m,ds];ref=group[REFERENCE,ds]
            summary.append([m,names[m],ds,g['n'],g['auc'],g['auc_sd'],g['aupr'],g['aupr_sd'],
                            g['auc']-ref['auc'],g['aupr']-ref['aupr']])
    header=['模块 / 变体','MDR AUC','MDR AUPR','ΔAUC vs 参考','ΔAUPR vs 参考',
            'MDS AUC','MDS AUPR','ΔAUC vs 参考','ΔAUPR vs 参考']
    own_header=['模块 / 变体','MiDrug AUC','MiDrug AUPR','ΔAUC vs 参考','ΔAUPR vs 参考']
    def table(selected):
        result=[]
        for m in order:
            row=[names[m]]
            for ds in selected:
                g=group[m,ds];ref=group[REFERENCE,ds]
                row += [g['auc'],g['aupr']]
                row += ['—','—'] if m==REFERENCE else [g['auc']-ref['auc'],g['aupr']-ref['aupr']]
            result.append(row)
        return result
    mainrows=table(('MDR','MDS'));ownrows=table(('MiDrug',))
    write_csv(HERE/'MDR_MDS_基线对比.csv',header,mainrows)
    write_csv(HERE/'MiDrug_基线对比.csv',own_header,ownrows)
    sh=['model_id','模块 / 变体','dataset','n','AUC mean','AUC SD','AUPR mean','AUPR SD','ΔAUC vs iNcRD-HG','ΔAUPR vs iNcRD-HG']
    write_csv(HERE/'summary.csv',sh,summary)
    indexed={(r['model'],r['dataset'],r['fold'],r['seed']):r for r in raw}
    detail=[]
    for r in raw:
        ref=indexed[REFERENCE,r['dataset'],r['fold'],r['seed']]
        detail.append([r['model'],r['dataset'],r['fold'],r['seed'],r['auc'],r['aupr'],
            r['auc']-ref['auc'],r['aupr']-ref['aupr'],r['best_epoch'],r['epochs_run'],r['val_auc'],
            r['elapsed_seconds'],r['run_origin'],r['split_sha256'],r['result_path'],r['checkpoint_sha256'],
            json.dumps(r['config'],ensure_ascii=False)])
    dh=['model_id','dataset','fold','seed','AUC','AUPR','配对ΔAUC','配对ΔAUPR','best_epoch','epochs_run',
        'validation_AUC','training_seconds','run_origin','split_sha256','result_path','checkpoint_sha256','config_JSON']
    write_csv(HERE/'per_run.csv',dh,detail)
    notes=[
        ('完成范围','14 个模型 × 3 数据集 = 42 个完整组合；182 次正式运行。169 次为 2026-09-28 结果，本次逐预测重算和哈希复核；13 次 M1+M2+M3 为 2026-10-04 新运行，均完成 300 epochs。'),
        ('参考及差值','所有 Δ = 当前模型实测均值 − 同数据集 iNcRD-HG 实测均值，单位为 AUC/AP 原始量纲，不是百分数。Excel/CSV 保存完整精度，显示四位小数；使用未舍入均值计算 Δ。'),
        ('MiDrug','iNcRD-HG miRNA–drug 发布划分；605 miRNA、216 药物。保留 train/val/test 成员，训练集同标签重复行 2 条去重。负正比 2:1。固定划分、3 个初始化种子 42/43/44。'),
        ('MDR / MDS','PLMF-MDA 的耐药 / 敏感关联数据。按项目固定 seed=42 外层 5 折，每折训练部分再分层留出 10% 验证。平衡正负 1:1。作者公开代码是 80/20，本项目的 5 折索引不是作者发布索引。'),
        ('统计','主表为 AUC / AP 的运行均值；均值与标准差表提供样本 SD(ddof=1)。MiDrug SD 反映初始化波动，MDR/MDS SD 反映一套 5 折的折间波动；不当作独立重复，不据此宣称显著性。'),
        ('验证选轮','所有方法只用验证集 AUC 选择 checkpoint，训练曲线没有逐轮测试指标；选定后才计算测试预测。没有以测试 AUC/AUPR 选模型或超参数。'),
        ('AUPR 定义','sklearn.metrics.average_precision_score，即 AP；不是梯形 PR 曲线面积。MiDrug 与 MDR/MDS 的正负比例不同，不跨数据集比较 AP 高低。'),
        ('图边界','所有目标关联图、GIP 和结构统计只从共享训练正例构建；验证和测试正例不加入消息图。训练、验证、测试均使用这张固定训练图。'),
        ('iNcRD-HG','原始代码提交 eb8de624a731a4c076143bf2f9c91e539affee43：GeneralConv×2、hidden=256、关系 mean、MLP，RMSprop(lr=1e-4, alpha=.99)，300 epochs。'),
        ('MiDrug 辅助图','iNcRD-HG 和 iNcRD-HG + M1 + M2 + M3 均保留原始 lncRNA/miRNA/drug 三类节点、9 类关系以及发布表达谱/ChemBERTa 特征；去掉后来添加的 circRNA。与部分序列模型的可用输入不同，属于各方法原生/适配输入比较。'),
        ('PLMF 跨数据适配','MDR/MDS 没有发布表达谱或 lncRNA 辅助关系：两个 iNcRD-HG 行均使用 80 维 2+3-mer、768 维原版 ChemBERTa、双向 kNN(10) 图、4 类关系。因此这两个数据集上的 iNcRD-HG 为数据适配版。'),
        ('M1','QMP：查询边对两层消息传递的精确修正。本次将原两类节点实现中的 1/2 改为实际目标节点的入关系均值系数，并包括所有 lin_self 项；保留原始辅助图。无新增编码器权重。'),
        ('M2','SP-RNCN：原实现的 co/knn/dknn 配对汇聚、扩展路径统计，ncn_dim=64；传播至中间节点的修正同样按目标关系数缩放。'),
        ('M3','MC-EASE：原实现的两侧掩蔽单元 EASE，lambda=10,50，4 维协同分数；无 M4 序列关系、无 ease_side 外部核。'),
        ('模块验收','训练前和训练后各在每数据集选 6 个训练正例，与显式删除查询边重算比较：原始编码器、端点修正、M2 汇聚、M2 静态统计、M3 和最终 logit 均通过；另验证梯度有限、修改留出标签不会改变图或结构特征。'),
        ('原示例数值','用户示例 0.9242/0.9161、0.9288/0.9252 出自 plmf_nolqo_20260923/REPORT_NOLQO.md 的旧 plain 实现 5 折×3种子汇总。本次参考是原始源码适配版的同一套 5 折实测，不能用示例常数替换。'),
        ('旧 MiDrug 改进分数','own_midrug_pilot_20260928 的旧结果去掉 lncRNA、在测试图加入验证边，且种子设置不同。本次不采用这些旧分数；只复用经核验的 M2/M3 实现并重新训练。'),
        ('适配标签','DLMVF 用 k-mer 替代缺失基因视图；AMMGC/GCMDR/LAGCN 为本地重实现；SubMDTA 只有监督部分；其余迁移/适配范围见“模型来源”。这些标签随模型名保留。'),
        ('基准范围','常规转导式配对划分，不是冷启动。按发布 ID 无跨集合对重叠，但成熟序列+原始 SMILES 相同的内容对仍有少量跨集合重合，沿用原基准并记录于数据表。负例来自公开未关联池，不全是实验真阴性。'),
        ('历史调参范围','模块配置来自此前实验，未在本轮搜索。共享评测提高了表内可比性，但本轮不是全新盲测，也不是嵌套超参数评估。'),
    ]
    wb=Workbook();wb.remove(wb.active)
    sheets=[('MDR_MDS_按要求格式',header,mainrows),('MiDrug_按要求格式',own_header,ownrows),
            ('均值与标准差',sh,summary),('逐次运行与配对差值',dh,detail),('实验说明',['项目','说明'],notes)]
    rh=['id','name','category','year','journal','doi','code','implementation','rationale']
    sheets.append(('模型来源',rh,[[r.get(k,'') for k in rh] for r in registry]))
    data_manifest=json.loads((BENCH/'data/manifest.json').read_text())
    data_rows=[]
    for ds,info in data_manifest['datasets'].items():
        for fold,rr in info['folds'].items():
            d=load(ds,int(fold))
            data_rows.append([ds,info['n_mirna'],info['n_drug'],fold,len(d['train_y']),int(sum(d['train_y'])),
                len(d['val_y']),len(d['test_y']),int(sum(d['test_y'])),rr['content_pair_train_test_overlap'],rr['sha256']])
    sheets.append(('数据与划分',['dataset','miRNA数','药物数','fold','训练数','训练正例','验证数','测试数','测试正例','内容相同train-test对数','split_sha256'],data_rows))
    for title,cols,rows in sheets:
        ws=wb.create_sheet(title);ws.append(cols)
        for row in rows:ws.append(row)
        ws.freeze_panes='B2' if title.endswith('按要求格式') else 'A2'
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='254B70')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[1].height=34
        for i,col in enumerate(cols,1):
            width=34 if i==1 else 21
            if title=='实验说明':width=24 if i==1 else 115
            elif col in ('result_path','config_JSON','code','implementation','rationale'):width=65
            ws.column_dimensions[get_column_letter(i)].width=width
        for cells in ws.iter_rows(min_row=2):
            for c in cells:
                c.alignment=Alignment(vertical='center',wrap_text=True)
                if isinstance(c.value,float):c.number_format='0.0000'
            if cells[0].value=='iNcRD-HG + M1 + M2 + M3':
                for c in cells:c.fill=PatternFill('solid',fgColor='E2EFDA');c.font=Font(bold=True)
            if title=='实验说明':ws.row_dimensions[cells[0].row].height=62
        if title.endswith('按要求格式'):
            for i,col in enumerate(cols,1):
                if col.startswith('Δ'):
                    for row in range(2,ws.max_row+1):
                        c=ws.cell(row,i)
                        if isinstance(c.value,float):c.number_format='+0.0000;-0.0000;0.0000'
            for wsrow in range(2,ws.max_row+1):ws.row_dimensions[wsrow].height=32
    workbook=ROOT/'三数据集_14模型基线_M1M2M3_20261004.xlsx'
    wb.save(workbook)
    review=load_workbook(workbook,data_only=True)
    assert review['MDR_MDS_按要求格式'].max_row==15
    assert review['MiDrug_按要求格式'].max_row==15
    for i,row in enumerate(mainrows,2):
        for j,value in enumerate(row,1):
            actual=review['MDR_MDS_按要求格式'].cell(i,j).value
            if isinstance(value,float):assert abs(actual-value)<1e-12
            else:assert actual==value
    full=group['incrd_m123','MiDrug']
    report=['# 三数据集、14 模型基线对比（2026-10-04）','',
            '完成 182/182 次运行、42/42 个模型×数据集组合；169 次已有结果经重新核验，13 次完整模块模型为本次新训练。',
            '参考统一为每个数据集的原始 iNcRD-HG 实测均值；Δ = 当前模型 − 参考。数值显示四位小数。','',
            '## PLMF-MDA：MDR / MDS','',markdown(header,mainrows),'',
            '## iNcRD-HG：MiDrug','',markdown(own_header,ownrows),'','## 协议与来源','']
    report += [f'- **{k}**：{v}' for k,v in notes]
    report += ['', '## 核验与复现', '',
        '- `verification.json`：182 条结果、预测、指标、checkpoint、划分哈希和验证集选轮核验。',
        '- `module_checks_initial.json` / `module_checks_trained.json`：训练前后显式删边等价性。',
        '- `per_run.csv`：每折/每种子值、与同一单元参考的配对差值及原始文件位置。',
        '- `summary.csv`：完整精度均值、样本标准差与差值。',
        '- `provenance.json` / `source_snapshot/`：源码快照与运行来源。',
        '', '```bash', f'cd {HERE}', 'PY=/data1/licheng/miniconda3/envs/ncRNAdrug/bin/python',
        '$PY check_modules.py', '$PY launch.py  # per-run resume; fixed 13 jobs, two GPUs',
        '$PY check_modules.py --trained', '$PY summarize.py', '```', '',
        f'已有 13 个模型的训练入口和完整参数：`{BENCH}/README.md`，该目录的 `launch.py` 与 `results/*.json`。',
        '如需从零重新训练 M1+M2+M3，可对 run.py 指定新的 --output 目录，避免覆盖本次结果。','']
    (HERE/'REPORT.md').write_text('\n'.join(report),encoding='utf-8')
    snapshot=HERE/'source_snapshot';snapshot.mkdir(exist_ok=True)
    paths=list(HERE.glob('*.py'))+[SOURCE,BENCH/'bench_data.py',BENCH/'engine.py',BENCH/'incrd_adapter.py',
          BENCH/'graph_adapters.py',BENCH/'original_source/predict_association.py',BENCH/'registry.json',
          BENCH/'source_manifest.json',BENCH/'data/manifest.json',BENCH/'data/incrd_original_manifest.json']
    source_records=[]
    for p in paths:
        relative=p.relative_to(ROOT);dest=snapshot/relative
        dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
        source_records.append({'path':str(p),'sha256':sha(p),'snapshot':str(dest.relative_to(HERE))})
    provenance={'created_utc':datetime.now(timezone.utc).isoformat(),'python':sys.version,'workbook':str(workbook),
        'new_training_entry':str(HERE/'run.py'),'reused_baseline_directory':str(BENCH),
        'sources':source_records,'artifact_sha256':{str(p):sha(p) for p in [workbook,HERE/'REPORT.md',HERE/'summary.csv',HERE/'per_run.csv']}}
    (HERE/'provenance.json').write_text(json.dumps(provenance,indent=2))
    print(workbook)
    print(markdown(header,mainrows))
    print(markdown(own_header,ownrows))

if __name__=='__main__':main()
