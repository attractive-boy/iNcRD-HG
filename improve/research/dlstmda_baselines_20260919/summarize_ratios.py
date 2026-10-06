# -*- coding: utf-8 -*-
"""Summarise the 1:5 / 1:10 ratio reproductions and compare with Supp. Table 1."""
import csv
import os
import statistics as st

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(BASE, 'results')
OUT = os.path.join(BASE, 'DLST_MDA_ratios_baselines.xlsx')

PAPER = {
    5: {
        'DAESTB':   (0.8388, 0.6648, 0.7664, 0.9111, 0.9280, 0.9238),
        'GCNNMMA':  (0.8884, 0.5881, 0.6319, 0.9397, 0.9257, 0.6877),
        'SubMDTA':  (0.9065, 0.6528, 0.6771, 0.9424, 0.8936, 0.7007),
        'GraphDTA': (0.9052, 0.6534, 0.6594, 0.9479, 0.9350, 0.7702),
        'ML-DTI':   (0.9044, 0.6486, 0.7155, 0.9482, 0.9155, 0.7228),
    },
    10: {
        'DAESTB':   (0.8022, 0.6099, 0.6082, 0.9662, 0.9234, 0.9303),
        'GCNNMMA':  (0.9236, 0.4418, 0.3591, 0.9701, 0.9281, 0.5632),
        'SubMDTA':  (0.9397, 0.6071, 0.5779, 0.9659, 0.9026, 0.6380),
        'GraphDTA': (0.9355, 0.5569, 0.4885, 0.9702, 0.9362, 0.6635),
        'ML-DTI':   (0.9329, 0.6360, 0.5655, 0.9648, 0.9265, 0.7008),
    },
}
FILE = {
    'GraphDTA': 'ratio_1to{r}_graphdta.csv',
    'GCNNMMA': 'ratio_1to{r}_gcnnma.csv',
    'SubMDTA': 'ratio_1to{r}_submdta.csv',
    'DAESTB': 'ratio_1to{r}_daestb.csv',
    'ML-DTI': 'ratio_1to{r}_mldti.csv',
}


def mean(rows, k):
    return st.mean(float(x[k]) for x in rows)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ratio 1to5 1to10'
    hdr = ['比例', '方法', '复现 AUC', '论文 AUC', 'ΔAUC', '复现 AUPR', '论文 AUPR',
           'ΔAUPR', '复现 ACC', '论文 ACC', '复现 F1', '复现 MCC', '论文 MCC']
    ws.append(hdr)
    HF = Font(bold=True, color='FFFFFF')
    HFILL = PatternFill('solid', fgColor='305496')
    OUR = PatternFill('solid', fgColor='FFF2CC')
    POS = PatternFill('solid', fgColor='C6EFCE')
    NEG = PatternFill('solid', fgColor='FFC7CE')
    for c in ws[1]:
        c.font = HF
        c.fill = HFILL
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    r = 2
    print(f"{'ratio':6s} {'method':10s} {'repAUC':>7s} {'papAUC':>7s} {'dAUC':>7s} "
          f"{'repAUPR':>8s} {'papAUPR':>8s} {'dAUPR':>7s}")
    for ratio in (5, 10):
        for m in ['GraphDTA', 'GCNNMMA', 'SubMDTA', 'DAESTB', 'ML-DTI']:
            fp = os.path.join(RES, FILE[m].format(r=ratio))
            if not os.path.exists(fp):
                ws.append([f'1:{ratio}', m] + ['-'] * 11)
                r += 1
                continue
            rows = list(csv.DictReader(open(fp)))
            pa, pm, ps, psp, pauc, paupr = PAPER[ratio][m]
            ra, ru = mean(rows, 'auc'), mean(rows, 'aupr')
            ws.append([f'1:{ratio}', m, round(ra, 4), pauc, round(ra - pauc, 4),
                       round(ru, 4), paupr, round(ru - paupr, 4),
                       round(mean(rows, 'acc'), 4), pa, round(mean(rows, 'f1'), 4),
                       round(mean(rows, 'mcc'), 4), pm])
            for c in range(3, 14):
                ws.cell(r, c).alignment = Alignment(horizontal='center')
            ws.cell(r, 3).fill = OUR
            ws.cell(r, 6).fill = OUR
            if ra - pauc >= 0:
                ws.cell(r, 5).fill = POS
            else:
                ws.cell(r, 5).fill = NEG
            if ru - paupr >= 0:
                ws.cell(r, 8).fill = POS
            else:
                ws.cell(r, 8).fill = NEG
            print(f'1:{ratio:<4d} {m:10s} {ra:7.4f} {pauc:7.4f} {ra-pauc:7.4f} '
                  f'{ru:8.4f} {paupr:8.4f} {ru-paupr:7.4f}')
            r += 1
    for i, w in enumerate([7, 12, 10, 10, 8, 11, 11, 8, 10, 10, 10, 10, 10], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
