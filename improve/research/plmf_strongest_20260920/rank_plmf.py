# -*- coding: utf-8 -*-
"""Rank iNcRD-HG's strongest portable model against PLMF-MDA and its 4 reproduced
baselines on the MDR/MDS benchmark (5cv and released single-split protocols)."""
import glob
import os
import statistics as st

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)
BL = os.path.join(BASE, 'plmf_baselines_20260920', 'results')
OUT = os.path.join(HERE, 'PLMF_MDA_strongest_ranking.xlsx')

PAPER = {
    ('MDR', 'PLMF-MDA'): (0.9227, 0.9076), ('MDR', 'GCNNMMA'): (0.9070, 0.8736),
    ('MDR', 'SubMDTA'): (0.8936, 0.8514), ('MDR', 'GraphDTA'): (0.8086, 0.7979),
    ('MDR', 'ML-DTI'): (0.7660, 0.7870),
    ('MDS', 'PLMF-MDA'): (0.9301, 0.9207), ('MDS', 'GCNNMMA'): (0.9143, 0.8889),
    ('MDS', 'SubMDTA'): (0.8858, 0.8587), ('MDS', 'GraphDTA'): (0.7864, 0.7911),
    ('MDS', 'ML-DTI'): (0.7820, 0.7881),
}
BASELINES = ['PLMF-MDA', 'GCNNMMA', 'SubMDTA', 'GraphDTA', 'ML-DTI']

HEAD = Font(bold=True, color='FFFFFF', size=10)
HEADF = PatternFill('solid', fgColor='305496')
OURF = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='BFBFBF')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center')


def baseline_agg(dataset, split, method):
    frames = [pd.read_csv(f) for f in glob.glob(os.path.join(BL, '*.csv'))
              if not os.path.basename(f).startswith('_')]
    rows = pd.concat(frames, ignore_index=True)
    g = rows[(rows.dataset == dataset) & (rows.split == split) & (rows.method == method)]
    if not len(g):
        return None
    return (st.mean(g.auc), st.stdev(g.auc) if len(g) > 1 else 0.0,
            st.mean(g.aupr), st.mean(g.acc), len(g))


def strongest_agg(dataset, split, which='fused'):
    frames = [pd.read_csv(f) for f in glob.glob(os.path.join(HERE, 'results',
              f'strongest_{dataset}_{split}_seed*.csv'))]
    rows = pd.concat(frames, ignore_index=True)
    if which == 'ref':
        return (st.mean(rows.ref_auc), st.stdev(rows.ref_auc) if len(rows) > 1 else 0.0,
                st.mean(rows.ref_aupr), st.mean(rows.acc), len(rows))
    return (st.mean(rows.fused_auc), st.stdev(rows.fused_auc) if len(rows) > 1 else 0.0,
            st.mean(rows.fused_aupr), st.mean(rows.acc), len(rows))


def board(dataset, split):
    entries = []
    a, sd, u, acc, n = strongest_agg(dataset, split, 'ref')
    entries.append(dict(method='iNcRD-HG-strongest (单模型 TabICLv2)', auc=a, auc_sd=sd, aupr=u,
                        acc=acc, n=n, source='frozen TabICLv2 图上下文参考（无融合）', ours=True))
    a, sd, u, acc, n = strongest_agg(dataset, split, 'fused')
    entries.append(dict(method='iNcRD-HG-strongest (+anchor 融合)', auc=a, auc_sd=sd, aupr=u,
                        acc=acc, n=n, source='单模型 + anchor-shrinkage (w=0.3)', ours=True))
    for m in BASELINES:
        r = baseline_agg(dataset, split, m)
        if r is None:
            continue
        a, sd, u, acc, n = r
        entries.append(dict(method=m, auc=a, auc_sd=sd, aupr=u, acc=acc, n=n,
                            source='PLMF-MDA 论文方法' if m == 'PLMF-MDA' else '基线复现',
                            ours=False))
    entries.sort(key=lambda e: (-e['auc'], -e['aupr']))
    for i, e in enumerate(entries, 1):
        e['rank'] = i
        e['paper'] = PAPER.get((dataset, e['method']), (None, None))
    return entries


def write_sheet(wb, name, title, entries, split_label):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws.cell(1, 1).font = Font(bold=True, size=11)
    hdr = ['排名', '方法', 'AUC', 'AUC±', 'AUPR', 'ACC', '论文AUC', '论文AUPR', '运行数', '来源/备注']
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(2, c)
        cell.font = HEAD
        cell.fill = HEADF
        cell.alignment = CEN
        cell.border = BORD
    ws.freeze_panes = 'A3'
    for e in entries:
        r = ws.max_row + 1
        vals = [e['rank'], e['method'], e['auc'], e['auc_sd'], e['aupr'], e['acc'],
                e['paper'][0], e['paper'][1], e['n'], e['source']]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            if isinstance(v, float):
                cell.value = round(v, 4)
                cell.number_format = '0.0000'
            else:
                cell.value = v
            cell.alignment = LEFT if c in (2, 10) else CEN
            cell.border = BORD
            if e['ours']:
                cell.fill = OURF
                cell.font = Font(bold=True, size=10)
    for i, w in enumerate([6, 28, 9, 8, 9, 8, 9, 9, 7, 34], 1):
        ws.column_dimensions[chr(64 + i)].width = w


def main():
    wb = Workbook()
    ws0 = wb.active
    ws0.title = '说明'
    notes = [
        ('PLMF-MDA benchmark：本文最强模型排名', ''),
        ('本文方法', 'iNcRD-HG 最强框架的可移植子集：frozen TabICLv2 图上下文参考 + anchor-shrinkage 后融合'),
        ('移植范围', '仅图上下文 33 列（关联图 + k-mer/Morgan 相似度）；证据档/DepMap 档因 MDR/MDS 无数据未启用'),
        ('数据', 'PLMF-MDA MDRdataset / MDSdataset；5cv = 论文 5 折(seed42)，single = 官方 80/20(seed0-4)'),
        ('对照', 'PLMF-MDA 与 GCNNMMA/SubMDTA/GraphDTA/ML-DTI 的复现值（见 ../plmf_baselines_20260920）'),
        ('论文值', '论文 Fig.3 图例值，仅供参考'),
    ]
    for i, (a, b) in enumerate(notes, 1):
        ws0.cell(i, 1, a).font = Font(bold=True, size=10)
        ws0.cell(i, 1).alignment = LEFT
        ws0.cell(i, 2, b).alignment = LEFT
    ws0.column_dimensions['A'].width = 14
    ws0.column_dimensions['B'].width = 95

    for ds in ('MDR', 'MDS'):
        for split in ('5cv', 'single'):
            entries = board(ds, split)
            label = '论文 5 折 (seed42)' if split == '5cv' else '官方 80/20 (seed 0-4)'
            write_sheet(wb, f'{ds}_{split}', f'{ds}dataset · {label}：排名', entries, label)
            print(f'\n--- {ds} {split} ---')
            for e in entries:
                print(f"{e['rank']:>2} {e['method']:30s} AUC {e['auc']:.4f}  AUPR {e['aupr']:.4f}")
    wb.save(OUT)
    print('\nwrote', OUT)


if __name__ == '__main__':
    main()
