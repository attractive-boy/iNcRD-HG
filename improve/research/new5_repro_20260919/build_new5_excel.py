import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "五模型复现汇总.xlsx")

rows = [
    # method, journal, code, protocol, auc, aupr, acc, f1, mcc, paper_auc, paper_aupr, note
    ["MPHGNN", "BIB 2026", "官方代码", "自有 MiDrug benchmark, 5折",
     0.9755, 0.9415, 0.9246, 0.8866, 0.8301, 0.9780, 0.9370, "seed42; seed422 为 0.9743/0.9382"],
    ["DLST-MDA", "JCIM 2025", "官方代码", "1578x156 benchmark, 5折 (正负1:1)",
     0.9437, 0.9356, 0.8753, 0.8827, None, 0.9411, 0.9326, "论文表 Acc0.8736/MCC0.7532"],
    ["PLMF-MDA (MDR)", "PLoS CB 2026", "官方代码", "自有 MDR benchmark, 80/20 x 5 seed",
     0.9260, 0.9135, None, None, None, 0.9222, 0.9062, "5 seed 均值±std: 0.9260±0.0061 / 0.9135±0.0064"],
    ["PLMF-MDA (MDS)", "PLoS CB 2026", "官方代码", "自有 MDS benchmark, 80/20 x 5 seed",
     0.9255, 0.9145, None, None, None, 0.9301, 0.9207, "5 seed 均值±std: 0.9255±0.0014 / 0.9145±0.0034"],
    ["DLMVF", "JCIM 2026", "官方代码", "1578x156 benchmark, 5折 (正负1:1)",
     0.9591, 0.9514, 0.8970, 0.8989, None, 0.9611, 0.9543, "修复 datasetindex 别名/drug_offset/缺失邻接"],
    ["MSFFMDI", "Interdiscip. Sci. 2025", "无(重实现)", "1578x156 benchmark, 5折 (正负1:1)",
     0.9483, 0.9391, 0.8394, 0.8205, 0.6938, None, None, "按摘要重实现, 非数值复现; AUC/AUPR 为 mean±std 0.0044/0.0071"],
]

detail = {
    "PLMF-MDA (MDR)": [(0, 0.9227, 0.9084), (1, 0.9245, 0.9149), (2, 0.9308, 0.9187),
                       (3, 0.9347, 0.9215), (4, 0.9175, 0.9042)],
    "PLMF-MDA (MDS)": [(0, 0.9233, 0.9131), (1, 0.9255, 0.9094), (2, 0.9276, 0.9134),
                       (3, 0.9256, 0.9179), (4, 0.9256, 0.9187)],
    "MSFFMDI": [(0, 0.9496, 0.9405), (1, 0.9538, 0.9480), (2, 0.9416, 0.9285),
                (3, 0.9478, 0.9411), (4, 0.9485, 0.9376)],
}

wb = Workbook()
ws = wb.active
ws.title = "复现汇总"
hdr = ["方法", "期刊", "代码", "数据/协议", "复现 AUC", "复现 AUPR", "ACC", "F1", "MCC",
       "论文 AUC", "论文 AUPR", "备注"]
ws.append(hdr)
for r in rows:
    ws.append(r)

head_fill = PatternFill("solid", fgColor="DDEBF7")
red = PatternFill("solid", fgColor="FFC7CE")
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = head_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(hdr)):
    for c in row:
        c.alignment = Alignment(vertical="center", wrap_text=True)
# highlight where paper > reproduction (AUC col idx 10, AUPR 11 vs rep 5,6)
for i, r in enumerate(rows, start=2):
    if r[9] is not None and r[4] is not None and r[9] > r[4]:
        ws.cell(i, 10).fill = red
    if r[10] is not None and r[5] is not None and r[10] > r[5]:
        ws.cell(i, 11).fill = red
for col, w in zip("ABCDEFGHIJKL", [18, 22, 12, 34, 10, 10, 9, 9, 9, 10, 10, 46]):
    ws.column_dimensions[col].width = w

ws2 = wb.create_sheet("逐seed-fold")
ws2.append(["方法", "seed/fold", "AUC", "AUPR"])
for m, vals in detail.items():
    for s, a, p in vals:
        ws2.append([m, s, a, p])
for c in ws2[1]:
    c.font = Font(bold=True)
    c.fill = head_fill
ws2.column_dimensions["A"].width = 20

wb.save(OUT)
print("wrote", OUT)
