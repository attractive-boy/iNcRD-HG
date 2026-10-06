# -*- coding: utf-8 -*-
"""MPHGNN_on_PLMF-MDA.xlsx: 说明 / 4 个榜单 (MDR_5cv, MDR_single, MDS_5cv, MDS_single) / 逐次运行."""
import glob, os, json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import summarize as S

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'MPHGNN_on_PLMF-MDA.xlsx')
HEAD = Font(bold=True, color='FFFFFF', size=10); HEADF = PatternFill('solid', fgColor='305496')
MPF = PatternFill('solid', fgColor='DDEBF7'); OURF = PatternFill('solid', fgColor='FFF2CC'); PAPF = PatternFill('solid', fgColor='E2EFDA')
thin = Side(style='thin', color='BFBFBF'); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True); LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def hdr(ws, row, cols):
    for c, v in enumerate(cols, 1):
        cell = ws.cell(row, c, v); cell.font = HEAD; cell.fill = HEADF; cell.alignment = CEN; cell.border = BORD


def main():
    wb = Workbook(); ws0 = wb.active; ws0.title = '说明'
    g = {ds: json.loads(str(dict(__import__('numpy').load(os.path.join(HERE, 'cache', f'graph_{ds}.npz')))['summary'])) for ds in ('MDR', 'MDS')}
    notes = [
        ('实验', '把 MPHGNN（Huang et al., Briefings in Bioinformatics 2026, 元路径引导异质图 GNN, miRNA-药物耐药）放到 PLMF-MDA（PLoS Comput Biol 2026）自建的 MDRdataset / MDSdataset 上跑，与 PLMF-MDA、其 4 个复现基线及本文最强模型同榜'),
        ('模型', '官方 model.py / layer.py 原样（github.com/Gtaizu/MPHGNN），仅把写死的 605 miRNA 偏移改成参数；官方超参：hidden 128、3 层 HeteroGCN、BiTrans 聚合、MIL+top-3 实例预测、k=20（441 条元路径实例/对）、batch 256、Adam lr 1e-4、pos_weight BCE'),
        ('异质图', f"MDR: {g['MDR']['n_mirna']} miRNA / {g['MDR']['n_drug']} drug / {g['MDR']['n_gene']} gene；miRNA-gene {g['MDR']['mirna_gene_edges']} 条（{g['MDR']['mirna_with_gene']} 个 miRNA 有靶基因），gene-drug {g['MDR']['gene_drug_edges']} 条（{g['MDR']['drug_with_gene']} 个药有靶点），gene-gene top-10 {g['MDR']['gene_gene_edges_topk']} 条。"
                  f"MDS: {g['MDS']['n_mirna']} / {g['MDS']['n_drug']} / {g['MDS']['n_gene']}；miRNA-gene {g['MDS']['mirna_gene_edges']}（{g['MDS']['mirna_with_gene']} miRNA），gene-drug {g['MDS']['gene_drug_edges']}（{g['MDS']['drug_with_gene']} drug），gene-gene {g['MDS']['gene_gene_edges_topk']}"),
        ('数据来源', 'miRNA-gene: ENCORI miRNATarget API（hg38, mRNA, clipExpNum>=5 且 >=3 个预测程序，每 miRNA 取 clip 支持最高的 50 个基因）；gene-drug: ChEMBL 机制靶点（DrugBank 本机被 403 拦截；SMILES→InChIKey→ChEMBL→含盐型子分子的 mechanism→UniProt→Entrez）；gene-gene: HumanNet v3 FN（LLS 每基因 top-10）；与 MPHGNN 原文同源（ENCORI / DrugBank→ChEMBL 替代 / HumanNet）'),
        ('节点特征', 'miRNA: 2+3-mer 频率 80 维（kmer 变体，全部节点可用）；或 CCLE 938 细胞系表达（expr 变体，仅 552/565 个 miRNA 有，其余置零+缺失标记）。drug: ChemBERTa-zinc-base-v1 mean-pool 768 维行内 min-max（与 MPHGNN 官方 drug.x 的构造方式逐位一致）。gene: HumanNet 对称归一化邻接的 200 维截断 SVD'),
        ('协议', '与 ../plmf_baselines_20260920 逐位相同的划分：5cv = 论文 5 折（正例+等量负例, seed 42）；single = 官方代码 80/20（seed 0-4）。每个 split 的消息传递 miRNA-drug 边只用训练正例；50 轮，记录测试集最优轮（AUC）与末轮、以及训练损失最小轮（官方模型选择口径）'),
        ('消融', 'gene-none：去掉全部基因边（无元路径，只剩二部图 + 节点特征）；gene-shuffled：随机置换基因 id（破坏 miRNA-gene / gene-drug 的生物学对应，边数不变）'),
        ('底色', '蓝=MPHGNN 变体；黄=本文最强模型（TabICLv2 图上下文 + anchor）；绿=PLMF-MDA 论文方法；白=复现基线'),
    ]
    for i, (a, b) in enumerate(notes, 1):
        ws0.cell(i, 1, a).font = Font(bold=True, size=10); ws0.cell(i, 1).alignment = LEFT; ws0.cell(i, 2, b).alignment = LEFT
        ws0.row_dimensions[i].height = 48
    ws0.column_dimensions['A'].width = 12; ws0.column_dimensions['B'].width = 120

    allrows = []
    for ds in ('MDR', 'MDS'):
        for split in ('5cv', 'single'):
            ws = wb.create_sheet(f'{ds}_{split}')
            ws.cell(1, 1, f"{ds}dataset · {'论文 5 折 (seed42)' if split == '5cv' else '官方 80/20 (seed 0-4)'}：排名（测试集最优轮）").font = Font(bold=True, size=11)
            cols = ['排名', '方法', 'AUC', 'AUC±', 'AUPR', 'ACC', '末轮AUC', '末轮AUPR', '最小loss轮AUC', '最小loss轮AUPR', '论文AUC', '论文AUPR', '运行数']
            hdr(ws, 2, cols); ws.freeze_panes = 'A3'
            for i, e in enumerate(S.board(ds, split), 1):
                p = e.get('paper') or (None, None)
                vals = [i, e['method'], e['auc'], e['auc_sd'], e['aupr'], e['acc'], e.get('last_auc'), e.get('last_aupr'),
                        e.get('minloss_auc'), e.get('minloss_aupr'), p[0], p[1], e['n']]
                r = ws.max_row + 1
                for c, v in enumerate(vals, 1):
                    cell = ws.cell(r, c)
                    if isinstance(v, float) and v == v:
                        cell.value = round(v, 4); cell.number_format = '0.0000'
                    elif v is None or (isinstance(v, float) and v != v):
                        cell.value = '-'
                    else:
                        cell.value = v
                    cell.alignment = LEFT if c == 2 else CEN; cell.border = BORD
                    cell.fill = {'mphgnn': MPF, 'ours': OURF, 'paper': PAPF}.get(e['kind'], PatternFill())
                    if e['kind'] in ('mphgnn', 'ours'):
                        cell.font = Font(bold=True, size=10)
            for i, w in enumerate([6, 60, 9, 8, 9, 8, 9, 9, 12, 12, 9, 9, 7], 1):
                ws.column_dimensions[chr(64 + i)].width = w
    ws = wb.create_sheet('逐次运行')
    fs = sorted(glob.glob(os.path.join(HERE, 'results', 'mphgnn_*.csv')))
    d = pd.concat([pd.read_csv(f) for f in fs], ignore_index=True)
    hdr(ws, 1, list(d.columns))
    for row in d.itertuples(index=False):
        ws.append([round(v, 4) if isinstance(v, float) else v for v in row])
    ws.freeze_panes = 'A2'
    wb.save(OUT); print('wrote', OUT)


if __name__ == '__main__':
    main()
