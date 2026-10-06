# -*- coding: utf-8 -*-
"""Coverage audit for candidate external relations on PLMF-MDA (MDR/MDS):
dz (CTD drug-disease x HMDD), gene (miRTarBase/ENCORI targets x LINCS/GEO/DrugMatrix/ChEMBL drug gene sets),
cerna (iNcRD-HG lncRNA-miRNA x lncRNA-drug), cluster (miRBase genomic clusters), pgx (GDSC2 IC50 profiles)."""
import os, sys, re, json, gzip, numpy as np, pandas as pd
sys.path.insert(0, '/data1/licheng/incrd-hg/iNcRD-HG/improve/research/mphgnn_plmf_20260921')
from plmf_splits import load_dataset
EXT = '/data1/licheng/incrd-hg/data_external'
DT = json.load(open('/data1/licheng/incrd-hg/iNcRD-HG/improve/research/mphgnn_plmf_20260921/cache/drug_targets.json'))['drugs']

def norm(s): return re.sub(r'[^a-z0-9]', '', str(s).lower())

for ds in ('MDR', 'MDS'):
    mir, dru, pos, neg = load_dataset(ds)
    names = mir['miRNA_Name'].astype(str).str.lower().values; dbs = dru['DrugBank_ID'].astype(str).values
    pref = {db: DT.get(db, {}).get('pref_name') for db in dbs}
    print(f'== {ds}: {len(names)} miRNAs, {len(dbs)} drugs; pref_name known for {sum(v is not None for v in pref.values())}')
    deg = np.bincount(pos[:, 0], minlength=len(names))
    # --- CTD drug-disease
    ctd = pd.read_csv(f'{EXT}/ncrna_disease_20260917/CTD_chemicals_diseases.tsv.gz', sep='\t', comment='#', header=None,
                      names=['chem', 'cid', 'cas', 'dis', 'did', 'ev', 'gene', 'score', 'omim', 'pmids'], usecols=[0, 1, 3, 4, 5], dtype=str)
    ctd = ctd[ctd['ev'].notna()]
    chem_norm = {norm(c): c for c in ctd['chem'].unique()}
    hit = {db: chem_norm.get(norm(p)) for db, p in pref.items() if p}
    n_hit = sum(v is not None for v in hit.values())
    print(f'CTD chem-disease: drugs matched by pref_name {n_hit}/{len(dbs)}; unmatched: {[pref[db] for db in dbs if hit.get(db) is None][:25]}')
    h = pd.read_csv(f'{EXT}/ncrna_disease_20260917/hmdd_v4_alldata.txt', sep='\t', encoding='latin1', on_bad_lines='skip')
    hd = set(h['disease'].astype(str).str.strip()); cd = set(ctd['dis'].unique())
    print(f'HMDD diseases {len(hd)}, CTD diseases {len(cd)}, exact-name overlap {len(hd & cd)}; HMDD assoc rows whose disease is in CTD: {h["disease"].astype(str).str.strip().isin(cd).mean():.2f}')
    # --- gene sets
    lincs = set(); geo_db = {}; dm = set(); cmap = set()
    for line in open(f'{EXT}/perturbation_20260917/LINCS_L1000_Chem_Pert_down.gmt'):
        t = line.split('\t')[0]; m = re.match(r'.*?-(.+)-[0-9.]+$', t.split(' ', 2)[-1] if ' ' in t else t)
        lincs.add(norm(t.split('H-', 1)[-1].rsplit('-', 1)[0]) if 'H-' in t else norm(t))
    for line in open(f'{EXT}/perturbation_20260917/Drug_Perturbations_from_GEO_up.gmt'):
        t = line.split('\t')[0]; m = re.search(r'(DB\d{5})', t)
        if m: geo_db.setdefault(m.group(1), 0); geo_db[m.group(1)] += 1
    for line in open(f'{EXT}/perturbation_20260917/DrugMatrix.gmt'):
        dm.add(norm(line.split('\t')[0].split('-')[0]))
    for line in open(f'{EXT}/ncrna_disease_20260917/perturbation_gmt/OldCMAP_up.gmt'):
        cmap.add(norm(line.split('\t')[0].rsplit('-', 1)[0]))
    def cov(S): return sum(1 for db in dbs if pref[db] and norm(pref[db]) in S)
    print(f'gene sets: LINCS name-match {cov(lincs)}/{len(dbs)}, GEO DBid-match {sum(db in geo_db for db in dbs)}, DrugMatrix {cov(dm)}, OldCMAP {cov(cmap)}, ChEMBL targets {sum(1 for db in dbs if DT.get(db, {}).get("genes"))}')
    mt = pd.read_csv(f'{EXT}/pathway_20260913/mirtarbase/hsa_MTI.csv', usecols=['miRNA', 'Target Gene', 'Support Type'])
    mtn = set(mt['miRNA'].astype(str).str.lower()); mts = set(mt[mt['Support Type'] == 'Functional MTI']['miRNA'].astype(str).str.lower())
    enc = pd.read_csv('/data1/licheng/incrd-hg/iNcRD-HG/improve/research/mphgnn_plmf_20260921/cache/encori_targets.tsv.gz', sep='\t', usecols=['miRNAname', 'geneName', 'clipExpNum'])
    en = set(enc['miRNAname'].astype(str).str.lower()); en3 = set(enc[enc['clipExpNum'] >= 3]['miRNAname'].astype(str).str.lower())
    inm = np.array([n in mtn for n in names]); ins = np.array([n in mts for n in names]); ine = np.array([n in en for n in names]); ine3 = np.array([n in en3 for n in names])
    print(f'miRNA targets: miRTarBase any {inm.mean():.2f} (deg<=1 miRNAs: {inm[deg<=1].mean():.2f}), functional {ins.mean():.2f}; ENCORI any {ine.mean():.2f} (deg<=1: {ine[deg<=1].mean():.2f}), clip>=3 {ine3.mean():.2f}')
    # --- ceRNA
    lm = pd.read_csv('/data1/licheng/incrd-hg/incrd-run/2.construct_network/LncMiRNA_interaction.csv'); ld = pd.read_csv('/data1/licheng/incrd-hg/incrd-run/2.construct_network/LncDrug_edge.csv')
    lmm = set(lm['miRNA'].astype(str).str.lower()); inl = np.array([n in lmm for n in names])
    print(f'ceRNA: miRNAs with lncRNA partners {inl.mean():.2f} (n={inl.sum()}, deg<=1: {inl[deg<=1].mean():.2f}); LncDrug drugs matched by DrugBank id {sum(db in set(ld["DrugBank_ID"].astype(str)) for db in dbs)}/{len(dbs)}; lncRNAs shared {len(set(lm["lncRNA"]) & set(ld["ncRNA_Name"]))}')
    # --- miRBase clusters
    rows = [l.rstrip('\n').split('\t') for l in open(f'{EXT}/mirbase_20260922/hsa.gff3') if not l.startswith('#')]
    pre = {}; mat = []
    for r in rows:
        attr = dict(kv.split('=') for kv in r[8].split(';'))
        if r[2] == 'miRNA_primary_transcript': pre[attr['ID']] = (r[0], int(r[3]), int(r[4]), r[6], attr['Name'])
        else: mat.append((attr['Name'].lower(), attr['Derives_from']))
    matn = {}
    for n, p in mat: matn.setdefault(n, []).append(p)
    inb = np.array([n in matn for n in names])
    print(f'miRBase v23.1: PLMF miRNA names found {inb.mean():.3f} (n missing {(~inb).sum()}): {[n for n, ok in zip(names, inb) if not ok][:12]}')
    # --- GDSC
    try:
        import openpyxl
        wb = openpyxl.load_workbook(f'{EXT}/drug_response_20260916/gdsc/GDSC2_fitted_dose_response_27Oct23.xlsx', read_only=True); ws = wb.active
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]; gn = set()
        for r in ws.iter_rows(min_row=2, values_only=True): gn.add(norm(r[hdr.index('DRUG_NAME')]))
        print(f'GDSC2 drugs {len(gn)}; PLMF matched by name {cov(gn)}/{len(dbs)}')
    except Exception as e:
        print('GDSC unreadable:', str(e)[:100])
